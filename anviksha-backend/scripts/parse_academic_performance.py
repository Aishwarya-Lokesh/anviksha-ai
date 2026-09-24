"""
Parses the two real spreadsheet formats found in the uploaded 'Academic Performance Batchwise'
folder:

Format A — "CGPA" sheet (used in Batch-2021, 2022, 2023, 2024 files):
    Row1: SL.NO | USN NO. | STUDENT NAMES | I (merged) | II (merged) | ... | VIII (merged)
    Row2:                                 | CREDITS SGPA GPA CGPA | CREDITS SGPA GPA CGPA | ...
    Data starts row 3.

Format B — "MASTER" sheet (used in 2020-21 Batch.xlsx):
    Row1: SL.NO | NAME | USN | BRANCH | EMAIL ID | PHONE NUMBER | COLLEGE | DATE OF BIRTH |
          10TH/SSLC PERCENTAGE | 12TH/2ND PUC PERCENTAGE | DIPLOMA PERCENTAGE |
          SGPA OF 1ST SEMESTER ... SGPA OF 8TH SEMESTER | CGPA UPTO CURRENT SEMESTER

Deliberate privacy decision: EMAIL, PHONE NUMBER, and DATE OF BIRTH columns are read but
NEVER carried into the normalized output or the database. The platform has no use for them
and they are sensitive PII outside the scope of an academic dashboard.
"""

import re
import openpyxl
from datetime import datetime

ROMAN_TO_SEM = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7, "VIII": 8}
CURRENT_YEAR = datetime.now().year
PASSOUT_OFFSET_YEARS = 4  # standard 4-year B.E. program


def usn_to_admission_year(usn: str):
    """4GW21IS002 -> 2021"""
    m = re.search(r"4GW(\d{2})IS", usn.upper()) if usn else None
    if not m:
        return None
    yy = int(m.group(1))
    return 2000 + yy


def status_for_admission_year(admission_year):
    if admission_year is None:
        return "unknown"
    passout_year = admission_year + PASSOUT_OFFSET_YEARS
    return "passed_out" if passout_year <= CURRENT_YEAR else "active"


def parse_cgpa_sheet(path, sheet_name, batch_label):
    """Format A parser — returns list of normalized student dicts.

    Column layout varies slightly between batch files (some end each semester block
    with a 'CGPA' column, others with 'Total GPA' and a separate trailing 'Final CGPA'
    column) — so columns are located by their row-2 label, not by a fixed offset.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]

    header_row = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    sub_row = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]

    # find each semester block's column range: [start, next_start)
    sem_starts = []  # (semester_number, start_idx 0-based)
    for idx, val in enumerate(header_row):
        if isinstance(val, str) and val.strip() in ROMAN_TO_SEM:
            sem_starts.append((ROMAN_TO_SEM[val.strip()], idx))
    sem_starts.sort(key=lambda x: x[1])

    blocks = []  # (semester_number, {label: col_idx_0based})
    for i, (sem_num, start) in enumerate(sem_starts):
        end = sem_starts[i + 1][1] if i + 1 < len(sem_starts) else len(header_row)
        label_map = {}
        for idx in range(start, end):
            label = sub_row[idx]
            if isinstance(label, str):
                label_map[label.strip().upper()] = idx
        blocks.append((sem_num, label_map))

    # a trailing "Final CGPA" column, if present, outside any semester block
    final_cgpa_idx = None
    for idx, label in enumerate(sub_row):
        if isinstance(label, str) and "FINAL CGPA" in label.strip().upper():
            final_cgpa_idx = idx

    students = []
    for row_idx in range(3, ws.max_row + 1):
        usn = ws.cell(row=row_idx, column=2).value
        name = ws.cell(row=row_idx, column=3).value
        if not usn or not name:
            continue
        usn = str(usn).strip()
        name = str(name).strip()

        semesters = {}
        last_real_cgpa = None
        for sem_num, label_map in blocks:
            sgpa_idx = label_map.get("SGPA")
            cgpa_idx = label_map.get("CGPA")
            sgpa = ws.cell(row=row_idx, column=sgpa_idx + 1).value if sgpa_idx is not None else None
            cgpa_after = ws.cell(row=row_idx, column=cgpa_idx + 1).value if cgpa_idx is not None else None

            if isinstance(cgpa_after, (int, float)) and 0 <= cgpa_after <= 10:
                last_real_cgpa = round(float(cgpa_after), 2)

            if sgpa is not None:
                semesters[sem_num] = {
                    "sgpa": round(float(sgpa), 2) if isinstance(sgpa, (int, float)) else None,
                    "cgpa_after_sem": round(float(cgpa_after), 2) if isinstance(cgpa_after, (int, float)) and 0 <= cgpa_after <= 10 else None,
                }

        # prefer an explicit trailing "Final CGPA" column if the sheet has one and it's a valid value
        final_val = ws.cell(row=row_idx, column=final_cgpa_idx + 1).value if final_cgpa_idx is not None else None
        if isinstance(final_val, (int, float)) and 0 <= final_val <= 10:
            final_cgpa = round(float(final_val), 2)
        else:
            final_cgpa = last_real_cgpa

        latest_sem = max(semesters.keys()) if semesters else None
        admission_year = usn_to_admission_year(usn)

        students.append({
            "usn": usn,
            "name": name,
            "batch": batch_label,
            "admission_year": admission_year,
            "status": status_for_admission_year(admission_year),
            "current_semester": latest_sem,
            "cgpa": final_cgpa,
            "semester_sgpa": semesters,
            "source_file": path,
            "source_sheet": sheet_name,
        })

    return students


def parse_master_sheet(path, sheet_name, batch_label):
    """Format B parser — returns list of normalized student dicts. Skips email/phone/DOB."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]

    headers = [str(ws.cell(row=1, column=c).value or "").strip().upper() for c in range(1, ws.max_column + 1)]

    def col_idx(*names):
        for n in names:
            for i, h in enumerate(headers):
                if n in h:
                    return i + 1
        return None

    name_col = col_idx("NAME")
    usn_col = col_idx("USN")
    cgpa_col = col_idx("CGPA UPTO CURRENT SEMESTER", "CGPA")
    sem_cols = {}
    for i, h in enumerate(headers):
        m = re.search(r"SGPA OF (\d)", h)
        if m:
            sem_cols[int(m.group(1))] = i + 1

    students = []
    for row_idx in range(2, ws.max_row + 1):
        usn = ws.cell(row=row_idx, column=usn_col).value if usn_col else None
        name = ws.cell(row=row_idx, column=name_col).value if name_col else None
        if not usn or not name:
            continue
        usn = str(usn).strip()
        name = str(name).strip()

        semesters = {}
        for sem_num, c in sem_cols.items():
            val = ws.cell(row=row_idx, column=c).value
            if isinstance(val, (int, float)):
                semesters[sem_num] = {"sgpa": round(float(val), 2), "cgpa_after_sem": None}

        cgpa = ws.cell(row=row_idx, column=cgpa_col).value if cgpa_col else None
        admission_year = usn_to_admission_year(usn)

        students.append({
            "usn": usn,
            "name": name,
            "batch": batch_label,
            "admission_year": admission_year,
            "status": status_for_admission_year(admission_year),
            "current_semester": max(semesters.keys()) if semesters else None,
            "cgpa": round(float(cgpa), 2) if isinstance(cgpa, (int, float)) else None,
            "semester_sgpa": semesters,
            "source_file": path,
            "source_sheet": sheet_name,
        })

    return students


def validate_student(rec):
    """Returns (is_complete: bool, issues: list[str])"""
    issues = []
    if rec.get("cgpa") is not None:
        if not (0 <= rec["cgpa"] <= 10):
            issues.append(f"CGPA out of range: {rec['cgpa']}")
    else:
        issues.append("Missing CGPA")
    if not rec.get("current_semester"):
        issues.append("Missing semester data")
    if not re.match(r"^4GW\d{2}IS\d{3}$", rec.get("usn", ""), re.IGNORECASE):
        issues.append("USN format unexpected")
    return (len(issues) == 0, issues)
