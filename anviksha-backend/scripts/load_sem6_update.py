"""
Loads real 6th semester SGPA data for the 2023-admission batch (currently in
their 7th semester) — this is genuinely new data, our records only had
semesters 1-5 for this batch before.

Only ADDS semester 6 to each student's semester_sgpa history — does NOT
touch semesters 1-5 already on file, even though this source file's own
5th-semester values differ slightly from what we already have (a real,
minor discrepancy between two different institutional snapshots, not
something to silently overwrite).

CGPA is updated to this file's own "CGPA (Upto Current Semester)" column —
using the institution's own authoritative computed value rather than
recomputing it ourselves, since VTU CGPA can involve credit-weighting we
don't want to guess at.

Usage:
    python scripts/load_sem6_update.py
"""

import sys
import os
import json

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import openpyxl
from app.core.database import SessionLocal
from app.models.models import Student, Department

DATA_PATH = os.path.join(
    os.path.dirname(__file__), "..", "data_raw", "sem6_update",
    "1786467380504_ISE_PLACEMENTS_DATABASE_7TH_SEM_2027_BATCH_.xlsx",
)


def main():
    db = SessionLocal()
    wb = openpyxl.load_workbook(DATA_PATH, data_only=True)
    ws = wb["Form Responses 1"]

    updated, unmatched, discrepancies = 0, [], []

    # Arundhati Das (4GW22IS012) — confirmed real student by the user: originally
    # admitted 2022, discontinued, rejoined and is now academically with the 2023
    # cohort. Her USN keeps the 2022 admission-year prefix even though she's
    # grouped with batch 2023 for all practical/current purposes.
    KNOWN_REJOINED_STUDENTS = {"4GW22IS012"}

    for r in range(2, ws.max_row + 1):
        usn = ws.cell(row=r, column=2).value
        if not usn:
            continue
        usn = str(usn).strip().upper()
        sem6_sgpa = ws.cell(row=r, column=16).value  # "6th Sem SGPA"
        sem5_sgpa_file = ws.cell(row=r, column=15).value  # for discrepancy check only
        cgpa = ws.cell(row=r, column=17).value  # "CGPA (Upto Current Semester)"
        backlogs = ws.cell(row=r, column=18).value

        student = db.query(Student).filter(Student.usn == usn).first()

        if not student and usn in KNOWN_REJOINED_STUDENTS:
            name = ws.cell(row=r, column=1).value.strip()
            dept = db.query(Department).first()
            history = {}
            for sem_i, col in [(1, 11), (2, 12), (3, 13), (4, 14), (5, 15), (6, 16)]:
                v = ws.cell(row=r, column=col).value
                if v is not None:
                    history[str(sem_i)] = {"sgpa": round(float(v), 2), "cgpa_after_sem": None}
            student = Student(
                usn=usn, name=name, batch="2023", department_id=dept.id,
                admission_year=2022, status="active", semester=7,
                cgpa=round(float(cgpa), 2) if cgpa else None,
                backlogs=int(backlogs) if backlogs else 0,
                data_status="verified", semester_sgpa=history,
            )
            db.add(student)
            db.flush()
            updated += 1
            continue

        if not student:
            unmatched.append(usn)
            continue

        history = dict(student.semester_sgpa) if student.semester_sgpa else {}

        # flag, don't silently overwrite, any mismatch in already-loaded semester 5
        existing_sem5 = history.get("5", {}).get("sgpa")
        if existing_sem5 is not None and sem5_sgpa_file is not None and abs(existing_sem5 - sem5_sgpa_file) > 0.01:
            discrepancies.append((usn, existing_sem5, sem5_sgpa_file))

        if sem6_sgpa is not None:
            history["6"] = {"sgpa": round(float(sem6_sgpa), 2), "cgpa_after_sem": round(float(cgpa), 2) if cgpa else None}
            student.semester_sgpa = history

        if cgpa is not None:
            student.cgpa = round(float(cgpa), 2)
        if backlogs is not None:
            student.backlogs = int(backlogs)
        if sem6_sgpa is not None:
            student.semester = 7  # 6th semester results in → now actively in 7th semester

        updated += 1

    db.commit()
    print(f"Updated {updated} students with real 6th semester SGPA + current CGPA.")
    if unmatched:
        print(f"{len(unmatched)} USNs from the file did not match any real student: {unmatched}")
    if discrepancies:
        print(f"\n{len(discrepancies)} students have a 5th-semester SGPA mismatch between our existing records and this file (not overwritten, just flagging):")
        for usn, old, new in discrepancies[:10]:
            print(f"  {usn}: existing={old}, this file={new}")
    db.close()


if __name__ == "__main__":
    main()
