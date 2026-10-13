"""Reads the material a student uploads for their project: the words, the
pictures and the tables in it. Nothing is interpreted here; this only makes
the student's own material available to the writing step."""
import csv
import io
import os
import re
import zipfile

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT = 120_000
MAX_IMAGES = 14
MAX_TABLES = 10
MAX_TABLE_ROWS = 30
MIN_IMAGE_SIDE = 140
ZIP_MAX_FILES = 80
ZIP_MAX_TOTAL = 60 * 1024 * 1024

DOCUMENTS = {".pdf", ".docx", ".pptx"}
TEXTS = {".txt", ".md", ".markdown", ".rst"}
IMAGES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
TABLES = {".csv", ".xlsx"}
CODE = {".py", ".js", ".jsx", ".ts", ".tsx", ".java", ".c", ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".php", ".rb", ".kt", ".swift", ".html", ".css",
        ".sql", ".sh", ".json", ".yml", ".yaml", ".xml", ".ipynb", ".r", ".m", ".dart", ".ino", ".toml", ".ini", ".cfg"}
ARCHIVES = {".zip"}
ALLOWED = DOCUMENTS | TEXTS | IMAGES | TABLES | CODE | ARCHIVES
SKIP_DIRS = ("node_modules/", "venv/", ".venv/", ".git/", "__pycache__/", "dist/", "build/", ".next/", "target/", ".idea/", ".vscode/")


class Unreadable(Exception):
    """The file is not something this feature can take, with a reason a student can act on."""


def kind_of(filename: str) -> str:
    ext = os.path.splitext(filename or "")[1].lower()
    for kind, group in (("document", DOCUMENTS), ("text", TEXTS), ("image", IMAGES), ("table", TABLES), ("code", CODE), ("archive", ARCHIVES)):
        if ext in group:
            return kind
    raise Unreadable(f"{ext or 'This kind of file'} can't be used. Upload PDF, Word, PowerPoint, text, images, CSV/Excel, code files or a .zip of code.")


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _image_ok(data: bytes) -> tuple[bool, int, int]:
    from PIL import Image
    try:
        with Image.open(io.BytesIO(data)) as im:
            im.verify()
        with Image.open(io.BytesIO(data)) as im:
            w, h = im.size
    except Exception:
        return False, 0, 0
    return (w >= MIN_IMAGE_SIDE and h >= MIN_IMAGE_SIDE and w * h <= 40_000_000), w, h


def as_png(data: bytes) -> bytes:
    """Any accepted picture as PNG or JPEG bytes that Word and PowerPoint both take."""
    from PIL import Image
    with Image.open(io.BytesIO(data)) as im:
        if im.format in ("PNG", "JPEG"):
            return data
        out = io.BytesIO()
        im.convert("RGB").save(out, format="PNG")
        return out.getvalue()


def _rows(table) -> list[list[str]]:
    rows = [[("" if c is None else str(c)).strip() for c in row] for row in table if row and any(c not in (None, "") for c in row)]
    return rows[:MAX_TABLE_ROWS]


def _pdf(data: bytes) -> dict:
    import pdfplumber
    from pypdf import PdfReader
    texts, tables, images = [], [], []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages[:150]:
                texts.append(page.extract_text() or "")
                if len(tables) < MAX_TABLES:
                    for t in page.extract_tables()[:2]:
                        rows = _rows(t)
                        if len(rows) >= 2 and len(rows[0]) >= 2:
                            tables.append({"caption": "", "rows": rows})
    except Exception as exc:
        raise Unreadable("This PDF could not be opened. If it is password-protected, remove the password and upload it again.") from exc
    try:
        for page in PdfReader(io.BytesIO(data)).pages[:80]:
            for image in page.images:
                if len(images) >= MAX_IMAGES:
                    break
                ok, _, _ = _image_ok(image.data)
                if ok:
                    images.append((image.name or f"image{len(images) + 1}.png", image.data))
    except Exception:
        pass                              # pictures are a bonus; the text is what matters
    text = _clean("\n\n".join(texts))
    note = None if len(text) > 80 else "No selectable text was found. If this is a scanned document, type the important points into the section questions."
    return {"text": text, "images": images, "tables": tables[:MAX_TABLES], "note": note}


def _docx(data: bytes) -> dict:
    from docx import Document
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise Unreadable("This Word file could not be opened. Save it again as .docx and upload that.") from exc
    text = _clean("\n".join(p.text for p in doc.paragraphs))
    tables = []
    for t in doc.tables[:MAX_TABLES]:
        rows = _rows([[c.text for c in r.cells] for r in t.rows])
        if len(rows) >= 2:
            tables.append({"caption": "", "rows": rows})
    images = []
    for rel in doc.part.rels.values():
        if "image" in rel.reltype and len(images) < MAX_IMAGES:
            try:
                blob = rel.target_part.blob
            except Exception:
                continue
            if _image_ok(blob)[0]:
                images.append((os.path.basename(str(rel.target_ref)), blob))
    return {"text": text, "images": images, "tables": tables, "note": None}


def _pptx(data: bytes) -> dict:
    from pptx import Presentation
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise Unreadable("This PowerPoint file could not be opened. Save it again as .pptx and upload that.") from exc
    lines, images = [], []
    for number, slide in enumerate(prs.slides, start=1):
        lines.append(f"[Slide {number}]")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
            if shape.shape_type == 13 and len(images) < MAX_IMAGES:
                try:
                    blob = shape.image.blob
                except Exception:
                    continue
                if _image_ok(blob)[0]:
                    images.append((f"slide{number}.{shape.image.ext}", blob))
    return {"text": _clean("\n".join(lines)), "images": images, "tables": [], "note": None}


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _table_file(filename: str, data: bytes) -> dict:
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:
            raise Unreadable("This Excel file could not be opened.") from exc
        rows = _rows(list(sheet.iter_rows(values_only=True))[: MAX_TABLE_ROWS + 5])
    else:
        rows = _rows(list(csv.reader(io.StringIO(_decode(data))))[: MAX_TABLE_ROWS + 5])
    text = "\n".join(" | ".join(r) for r in rows)
    return {"text": text, "images": [], "tables": [{"caption": os.path.splitext(os.path.basename(filename))[0].replace("_", " "), "rows": rows}] if len(rows) >= 2 else [], "note": None}


def _notebook(data: bytes) -> str:
    import json
    try:
        cells = json.loads(_decode(data)).get("cells", [])
    except Exception:
        return ""
    parts = []
    for cell in cells:
        source = "".join(cell.get("source", []))
        if source.strip():
            parts.append(("# [markdown]\n" if cell.get("cell_type") == "markdown" else "") + source)
    return "\n\n".join(parts)


def _zip(data: bytes) -> dict:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise Unreadable("This .zip file is damaged and could not be opened.") from exc
    infos = [i for i in archive.infolist() if not i.is_dir()]
    if sum(i.file_size for i in infos) > ZIP_MAX_TOTAL:
        raise Unreadable("This .zip unpacks to more than 60 MB. Upload only the source code, without libraries or data.")
    parts, images, used, listing = [], [], 0, []
    wanted = [i for i in infos if not any(d in i.filename.replace("\\", "/") + "/" for d in SKIP_DIRS)]
    # the README and the shortest paths first: they describe the project best
    wanted.sort(key=lambda i: (0 if "readme" in i.filename.lower() else 1, i.filename.count("/"), i.filename))
    for info in wanted[:ZIP_MAX_FILES]:
        ext = os.path.splitext(info.filename)[1].lower()
        listing.append(info.filename)
        if ext in IMAGES and len(images) < MAX_IMAGES and info.file_size <= 8 * 1024 * 1024:
            blob = archive.read(info)
            if _image_ok(blob)[0]:
                images.append((os.path.basename(info.filename), blob))
        elif (ext in CODE or ext in TEXTS) and info.file_size <= 400_000 and used < MAX_TEXT:
            body = _notebook(archive.read(info)) if ext == ".ipynb" else _decode(archive.read(info))
            body = body[:12_000]
            parts.append(f"===== {info.filename} =====\n{body}")
            used += len(body)
    text = "Files in the project:\n" + "\n".join(listing[:ZIP_MAX_FILES]) + "\n\n" + "\n\n".join(parts)
    return {"text": _clean(text), "images": images, "tables": [], "note": None if parts else "No readable source files were found in this .zip."}


def extract(filename: str, data: bytes) -> dict:
    """{"kind", "text", "images": [(name, bytes)], "tables": [{caption, rows}], "note"} for one uploaded file."""
    if not data:
        raise Unreadable("This file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise Unreadable("This file is larger than 20 MB. Upload a smaller version.")
    kind = kind_of(filename)
    ext = os.path.splitext(filename)[1].lower()
    if kind == "document":
        out = {".pdf": _pdf, ".docx": _docx, ".pptx": _pptx}[ext](data)
    elif kind == "image":
        ok, w, h = _image_ok(data)
        if not ok:
            raise Unreadable("This picture could not be read, or it is smaller than 140 pixels on a side.")
        out = {"text": "", "images": [(os.path.basename(filename), data)], "tables": [], "note": None}
    elif kind == "table":
        out = _table_file(filename, data)
    elif kind == "archive":
        out = _zip(data)
    else:
        body = _notebook(data) if ext == ".ipynb" else _decode(data)
        out = {"text": _clean((f"===== {os.path.basename(filename)} =====\n" if kind == "code" else "") + body), "images": [], "tables": [], "note": None}
    out["text"] = out["text"][:MAX_TEXT]
    out["kind"] = kind
    return out
