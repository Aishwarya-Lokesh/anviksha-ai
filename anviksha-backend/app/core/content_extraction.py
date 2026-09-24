"""Format-aware source extraction used by the PPT/notes generator."""

import logging
import re
from collections import Counter

import docx
from pypdf import PdfReader
from pptx import Presentation

logger = logging.getLogger(__name__)
EXPLICIT_HEADING_RE = re.compile(r"^\s*(?i:chapter|module|unit)\s+[\divxlcdm]+\s*(?:[:.\-]\s*|\s+)([A-Z][^\n]{2,120})$")
NUMBERED_HEADING_RE = re.compile(r"^\s*\d+(?:\.\d+){0,3}\.?\s+([A-Z][^\n]{2,120})$")
TOP_LEVEL_LABEL_RE = re.compile(r"^\s*(?i:chapter|module|unit)\s+[\divxlcdm]+\s*$")


def extract_ppt_style(filepath: str) -> dict:
    """Return structural statistics from a real .pptx style reference."""
    prs = Presentation(filepath)
    bullet_counts, heading_samples, words_per_bullet = [], [], []
    for slide in prs.slides:
        title_text, bullets_this_slide = None, 0
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                if shape == slide.shapes.title:
                    title_text = text
                else:
                    bullets_this_slide += 1
                    words_per_bullet.append(len(text.split()))
        if title_text:
            heading_samples.append(title_text)
        bullet_counts.append(bullets_this_slide)
    return {"source_type": "ppt", "slide_count": len(prs.slides), "avg_bullets_per_slide": round(sum(bullet_counts) / len(bullet_counts), 1) if bullet_counts else 0, "avg_words_per_bullet": round(sum(words_per_bullet) / len(words_per_bullet), 1) if words_per_bullet else 0, "heading_samples": heading_samples[:5]}


def extract_notes_style(filepath: str) -> dict:
    d = docx.Document(filepath)
    headings, lengths = [], []
    for para in d.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style.name.startswith("Heading"):
            headings.append(text)
        else:
            lengths.append(len(text.split()))
    return {"source_type": "notes", "heading_count": len(headings), "avg_paragraph_words": round(sum(lengths) / len(lengths), 1) if lengths else 0, "heading_samples": headings[:5]}


def _remove_repeated_margins(pages: list[dict]) -> list[dict]:
    """Remove only short lines repeated on 3+ pages (typical running headers)."""
    candidates = []
    for page in pages:
        lines = [line.strip() for line in page["text"].splitlines() if line.strip()]
        candidates.extend(set(lines[:3] + lines[-3:]))
    repeated = {line for line, count in Counter(candidates).items() if count >= 3 and len(line) <= 100}
    return [{**page, "text": "\n".join(line for line in page["text"].splitlines() if line.strip() not in repeated)} for page in pages]


_EXTRACTABLE_IMAGE_FILTERS = {"DCTDecode": "jpg", "JPXDecode": "jp2"}
_MIN_IMAGE_DIMENSION = 150  # skip small icons/bullets/decorative marks, real figures only


def extract_images_from_pdf(filepath: str) -> dict:
    """Real embedded figures pulled directly out of a real textbook PDF -
    never an AI-generated stand-in. Only JPEG/JPEG2000-filtered images are
    extracted (their raw PDF stream bytes ARE already a valid image file);
    other encodings (e.g. CCITT fax-compressed scans, common in scanned
    textbooks) aren't safely extractable this way and are honestly skipped
    rather than guessed at. Returns {page_number: {"data": bytes, "ext": str,
    "width": int, "height": int}} keeping only the largest real image per
    page when a page has more than one."""
    import pdfplumber

    images_by_page: dict = {}
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            for img in page.images:
                stream = img.get("stream")
                if stream is None:
                    continue
                filt = stream.get("Filter")
                filt_name = str(filt).strip("/'\"") if filt else None
                ext = _EXTRACTABLE_IMAGE_FILTERS.get(filt_name)
                if not ext:
                    continue
                width, height = img.get("width", 0), img.get("height", 0)
                if width < _MIN_IMAGE_DIMENSION or height < _MIN_IMAGE_DIMENSION:
                    continue
                try:
                    data = stream.rawdata
                except Exception:
                    continue
                existing = images_by_page.get(page.page_number)
                if existing is None or (width * height) > (existing["width"] * existing["height"]):
                    images_by_page[page.page_number] = {"data": data, "ext": ext, "width": int(width), "height": int(height)}
    return images_by_page


def extract_document_pages(filepath: str, filename: str) -> list[dict]:
    """Extract ordered text. A textless PDF is reported instead of faked."""
    ext = filename.lower().rsplit(".", 1)[-1]
    pages = []
    if ext == "pdf":
        pdf = PdfReader(filepath)
        for number, page in enumerate(pdf.pages, start=1):
            pages.append({"page": number, "text": (page.extract_text() or "").strip()})
        pages = _remove_repeated_margins(pages)
        extracted = sum(len(p["text"]) for p in pages)
        logger.info("PPT source PDF '%s': %d pages, %d extracted characters, %d words", filename, len(pages), extracted, sum(len(p["text"].split()) for p in pages))
        if not extracted:
            raise ValueError("This PDF contains no extractable text. It appears to be scanned/image-based; OCR is not configured on this server.")
        # A real body-text page in this course's PDFs runs well into the
        # hundreds of characters. A handful of stray characters per page
        # (e.g. a cover-page title rendered as real text over otherwise
        # scanned/image pages) is not enough real content to generate
        # anything meaningful from — surface that honestly now rather than
        # silently accepting a near-empty "chapter" that fails confusingly
        # deep inside AI generation later.
        if extracted / len(pages) < 80:
            raise ValueError(
                f"This PDF only contains {extracted} extractable characters across {len(pages)} page(s) "
                "(e.g. a cover title) — it appears to be a scanned/image-based document with no real "
                "body text. OCR is not configured on this server, so this file can't be used as a "
                "content source. Please upload a text-based PDF or a .docx version instead."
            )
    elif ext == "docx":
        d = docx.Document(filepath)
        text = "\n".join(p.text for p in d.paragraphs if p.text.strip())
        for table in d.tables:
            text += "\n" + "\n".join(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        pages = [{"page": 1, "text": text.strip()}]
    elif ext == "txt":
        with open(filepath, "r", encoding="utf-8", errors="replace") as source:
            pages = [{"page": 1, "text": source.read().strip()}]
    elif ext == "pptx":
        prs = Presentation(filepath)
        for number, slide in enumerate(prs.slides, start=1):
            pages.append({"page": number, "text": "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame and shape.text.strip())})
    else:
        raise ValueError(f"Unsupported content source format: {ext}")
    if not any(page["text"] for page in pages):
        raise ValueError(f"No readable text was found in '{filename}'.")
    return pages


def _heading_title(line: str) -> str | None:
    line = line.strip()
    match = EXPLICIT_HEADING_RE.match(line) or NUMBERED_HEADING_RE.match(line)
    if not match:
        return None
    # Sentences and running text are deliberately excluded: headings are short.
    number_count = len(re.findall(r"\b\d+(?:\.\d+){0,3}\b", line))
    return line if len(line) <= 140 and number_count <= 2 and not line.endswith((".", ",", ";")) else None


def _chunk_section(title: str, entries: list[tuple[int, str]], size: int = 6500) -> list[dict]:
    """Chunk by paragraphs, never arbitrary characters, retaining page provenance."""
    chunks, current, source_pages, current_length = [], [], [], 0
    for page, paragraph in entries:
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if current and current_length + len(paragraph) > size:
            chunks.append({"title": title, "text": "\n".join(current), "source_pages": source_pages})
            current, source_pages, current_length = [], [], 0
        current.append(paragraph)
        current_length += len(paragraph) + 1
        if page not in source_pages:
            source_pages.append(page)
    if current:
        chunks.append({"title": title, "text": "\n".join(current), "source_pages": source_pages})
    return chunks


def detect_chapters(filepath: str, filename: str) -> dict:
    """Detect headings while retaining their explanatory paragraphs and pages."""
    pages = extract_document_pages(filepath, filename)
    sections, current_title, current_entries = [], None, []
    for page in pages:
        page_heading_count = sum(1 for line in page["text"].splitlines() if _heading_title(line))
        if page_heading_count >= 5:  # table-of-contents page, not chapter prose
            continue
        lines = page["text"].splitlines()
        for line_index, line in enumerate(lines):
            # Top-level labels create selectable chapters; numbered subheadings
            # stay with their following paragraphs inside the chapter text.
            heading = line.strip() if EXPLICIT_HEADING_RE.match(line.strip()) else None
            if TOP_LEVEL_LABEL_RE.match(line.strip()) and line_index + 1 < len(lines):
                next_line = lines[line_index + 1].strip()
                if next_line and len(next_line) <= 120:
                    heading = f"{line.strip()}: {next_line}"
            if heading:
                if current_title and current_entries:
                    sections.extend(_chunk_section(current_title, current_entries))
                current_title, current_entries = heading, []
            elif line.strip():
                current_entries.append((page["page"], line.strip()))
    if current_title and current_entries:
        sections.extend(_chunk_section(current_title, current_entries))
    sections = [section for section in sections if len(section["text"]) >= 250]
    if not sections:
        sections = _chunk_section("Source content", [(page["page"], paragraph) for page in pages for paragraph in page["text"].split("\n") if paragraph.strip()])
    raw_chars = sum(len(page["text"]) for page in pages)
    logger.info("PPT source '%s': %d pages, %d characters, %d sections/chunks", filename, len(pages), raw_chars, len(sections))
    return {"chapters": sections, "chapters_detected": bool(current_title), "page_count": len(pages), "extracted_characters": raw_chars}
