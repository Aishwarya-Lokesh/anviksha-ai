"""Administration: login accounts, bulk student import, departments, subject
allocation, and placement / internship records. Every change is written to
the activity log."""
import csv
import io
import os
import re
import secrets
import string
import tempfile
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.core.activity_log import log_activity
from app.core.database import get_db
from app.core.placement_matching import canonical_company
from app.core.security import hash_password, require_role
from app.models.models import (
    Announcement, Company, Department, Faculty, Internship, Placement, PlacementDrive, Student, Subject, User,
)

router = APIRouter(prefix="/admin", tags=["admin-manage"])

ROLES = ("student", "teacher", "hod", "admin")
USN_RE = re.compile(r"^\d?[A-Z]{1,4}\d{2}[A-Z]{2}\d{2,4}$")
EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
OFFER_TYPES = ("full-time", "PPO", "internship+PPO")
OFFER_STATUSES = ("offered", "accepted", "joined", "declined")
MAX_IMPORT_ROWS = 2000


def _temp_password() -> str:
    alphabet = string.ascii_letters.replace("l", "").replace("I", "").replace("O", "") + "23456789"
    while True:
        pw = "".join(secrets.choice(alphabet) for _ in range(12))
        if any(c.isdigit() for c in pw) and any(c.isalpha() for c in pw):
            return pw


# ------------------------------------------------------------------ login accounts

def _login_out(db: Session, u: User) -> dict:
    name, detail = u.display_name, None
    if u.student_id:
        s = db.query(Student).filter(Student.id == u.student_id).first()
        if s:
            name, detail = s.name, s.usn
    elif u.faculty_id:
        f = db.query(Faculty).filter(Faculty.id == u.faculty_id).first()
        if f:
            name, detail = f.name, f.designation
    return {"id": u.id, "email": u.email, "role": u.role, "name": name, "detail": detail, "active": bool(u.is_active),
            "department_id": u.department_id}


@router.get("/logins")
def list_logins(role: Optional[str] = None, q: Optional[str] = None, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))):
    query = db.query(User)
    if role:
        query = query.filter(User.role == role)
    rows = [_login_out(db, u) for u in query.order_by(User.role, User.email).all()]
    if q:
        needle = q.strip().lower()
        rows = [r for r in rows if needle in (r["email"] or "").lower() or needle in (r["name"] or "").lower() or needle in (r["detail"] or "").lower()]
    return {"total": len(rows), "logins": rows[:300]}


class LoginCreate(BaseModel):
    email: str
    role: str
    student_id: Optional[int] = None
    faculty_id: Optional[int] = None
    department_id: Optional[int] = None
    display_name: Optional[str] = None


@router.post("/logins")
def create_login(payload: LoginCreate, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))):
    email = payload.email.strip().lower()
    if not EMAIL_RE.match(email):
        raise HTTPException(status_code=400, detail="That email address doesn't look right")
    if payload.role not in ROLES:
        raise HTTPException(status_code=400, detail="Role must be student, teacher, hod or admin")
    if db.query(User).filter(User.email == email).first():
        raise HTTPException(status_code=400, detail="A login with this email already exists")
    if payload.role == "student":
        if not payload.student_id or not db.query(Student).filter(Student.id == payload.student_id).first():
            raise HTTPException(status_code=400, detail="Choose the student this login belongs to")
        if db.query(User).filter(User.student_id == payload.student_id).first():
            raise HTTPException(status_code=400, detail="That student already has a login")
    if payload.role == "teacher":
        if not payload.faculty_id or not db.query(Faculty).filter(Faculty.id == payload.faculty_id).first():
            raise HTTPException(status_code=400, detail="Choose the faculty member this login belongs to")
        if db.query(User).filter(User.faculty_id == payload.faculty_id).first():
            raise HTTPException(status_code=400, detail="That faculty member already has a login")
    if payload.role == "hod" and (not payload.department_id or not db.query(Department).filter(Department.id == payload.department_id).first()):
        raise HTTPException(status_code=400, detail="Choose the department this HOD heads")

    password = _temp_password()
    u = User(email=email, hashed_password=hash_password(password), role=payload.role,
             student_id=payload.student_id if payload.role == "student" else None,
             faculty_id=payload.faculty_id if payload.role == "teacher" else None,
             department_id=payload.department_id if payload.role == "hod" else None,
             display_name=(payload.display_name or "").strip() or None)
    db.add(u)
    db.commit()
    db.refresh(u)
    log_activity(db, user, f"Created {payload.role} login {email}")
    # The password is shown once, here. It is stored only in scrambled form.
    return {**_login_out(db, u), "temporary_password": password}


class LoginEdit(BaseModel):
    role: Optional[str] = None
    active: Optional[bool] = None
    department_id: Optional[int] = None


def _login_or_404(db: Session, login_id: int) -> User:
    u = db.query(User).filter(User.id == login_id).first()
    if not u:
        raise HTTPException(status_code=404, detail="Login not found")
    return u


@router.patch("/logins/{login_id}")
def edit_login(login_id: int, payload: LoginEdit, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))):
    u = _login_or_404(db, login_id)
    sent = payload.dict(exclude_unset=True)
    if u.id == user.id and (sent.get("active") is False or ("role" in sent and sent["role"] != "admin")):
        raise HTTPException(status_code=400, detail="You can't deactivate your own account or remove your own admin role")
    changes = []
    if "role" in sent and sent["role"] != u.role:
        if sent["role"] not in ROLES:
            raise HTTPException(status_code=400, detail="Role must be student, teacher, hod or admin")
        if sent["role"] == "student" and not u.student_id:
            raise HTTPException(status_code=400, detail="This login isn't linked to a student record")
        if sent["role"] == "teacher" and not u.faculty_id:
            raise HTTPException(status_code=400, detail="This login isn't linked to a faculty record")
        if sent["role"] == "hod":
            dept_id = sent.get("department_id") or u.department_id
            if not dept_id and u.faculty_id:
                f = db.query(Faculty).filter(Faculty.id == u.faculty_id).first()
                dept_id = f.department_id if f else None
            if not dept_id:
                raise HTTPException(status_code=400, detail="Choose the department this HOD heads")
            u.department_id = dept_id
        changes.append(f"role {u.role} -> {sent['role']}")
        u.role = sent["role"]
    if "active" in sent and bool(sent["active"]) != bool(u.is_active):
        u.is_active = bool(sent["active"])
        changes.append("reactivated" if u.is_active else "deactivated")
    db.commit()
    if changes:
        log_activity(db, user, f"Login {u.email}: {', '.join(changes)}")
    return _login_out(db, u)


@router.post("/logins/{login_id}/reset-password")
def reset_password(login_id: int, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))):
    u = _login_or_404(db, login_id)
    password = _temp_password()
    u.hashed_password = hash_password(password)
    db.commit()
    log_activity(db, user, f"Reset the password of {u.email}")
    return {"email": u.email, "temporary_password": password}


# ------------------------------------------------------------------ bulk student import

def _read_table(filename: str, data: bytes) -> list[dict]:
    """Rows of a CSV or Excel file as dicts keyed by lower-cased header."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xls")):
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = [[("" if c is None else str(c)).strip() for c in row] for row in sheet.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if not rows:
            return []
        header = [h.lower() for h in rows[0]]
        return [dict(zip(header, r)) for r in rows[1:]]
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        return [{(k or "").strip().lower(): (v or "").strip() for k, v in row.items()} for row in csv.DictReader(io.StringIO(text))]
    raise HTTPException(status_code=400, detail="Upload a .csv or .xlsx file")


def _pick(row: dict, *names: str) -> str:
    for n in names:
        if row.get(n):
            return str(row[n]).strip()
    return ""


@router.post("/import/students")
async def import_students(
    file: UploadFile = File(...), apply: bool = Form(False), create_logins: bool = Form(False),
    department_id: Optional[int] = Form(None),
    db: Session = Depends(get_db), user: User = Depends(require_role("admin")),
):
    """Checks every row first and reports what would happen. Nothing is
    written unless `apply` is set, and then only the rows that passed."""
    rows = _read_table(file.filename or "", await file.read())
    if not rows:
        raise HTTPException(status_code=400, detail="That file has no rows")
    if len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(status_code=400, detail=f"Too many rows ({MAX_IMPORT_ROWS} at most per file)")
    dept = db.query(Department).filter(Department.id == department_id).first() if department_id else db.query(Department).first()
    if not dept:
        raise HTTPException(status_code=400, detail="Create a department first")

    existing = {u for (u,) in db.query(Student.usn).all()}
    seen: set[str] = set()
    report, ready = [], []
    for number, row in enumerate(rows, start=2):       # row 1 is the header
        usn, name = _pick(row, "usn", "roll no", "roll number").upper(), _pick(row, "name", "student name")
        batch, semester, cgpa = _pick(row, "batch", "admission year"), _pick(row, "semester", "sem"), _pick(row, "cgpa")
        errors = []
        if not USN_RE.match(usn):
            errors.append("USN is missing or not in the usual format (e.g. 4GW23IS001)")
        elif usn in existing:
            errors.append("a student with this USN already exists")
        elif usn in seen:
            errors.append("this USN appears more than once in the file")
        if len(name) < 2:
            errors.append("name is missing")
        if not re.match(r"^20\d{2}$", batch):
            errors.append("batch must be a four-digit entry year such as 2023")
        sem_value = None
        if semester:
            if semester.isdigit() and 1 <= int(semester) <= 8:
                sem_value = int(semester)
            else:
                errors.append("semester must be between 1 and 8")
        cgpa_value = None
        if cgpa:
            try:
                cgpa_value = float(cgpa)
                if not 0 <= cgpa_value <= 10:
                    raise ValueError
            except ValueError:
                errors.append("CGPA must be a number between 0 and 10")
        seen.add(usn)
        report.append({"row": number, "usn": usn, "name": name, "batch": batch, "ok": not errors, "problems": errors})
        if not errors:
            ready.append({"usn": usn, "name": name, "batch": batch, "semester": sem_value, "cgpa": cgpa_value, "section": _pick(row, "section") or None})

    created, logins = 0, []
    if apply and ready:
        for r in ready:
            s = Student(usn=r["usn"], name=r["name"], batch=r["batch"], admission_year=int(r["batch"]), status="active",
                        semester=r["semester"], section=r["section"], cgpa=r["cgpa"], backlogs=0, department_id=dept.id, data_status="provisional")
            db.add(s)
            db.flush()
            created += 1
            if create_logins:
                email = f"{r['usn'].lower()}@anviksha.ai"
                if not db.query(User).filter(User.email == email).first():
                    password = _temp_password()
                    db.add(User(email=email, hashed_password=hash_password(password), role="student", student_id=s.id))
                    logins.append({"usn": r["usn"], "email": email, "temporary_password": password})
        db.commit()
        log_activity(db, user, f"Imported {created} student(s) from {file.filename}")
    return {"department": dept.name, "rows": len(rows), "valid": len(ready), "invalid": len(rows) - len(ready),
            "applied": bool(apply), "created": created, "logins": logins, "report": report[:500]}


# ------------------------------------------------------------------ departments and subjects

class DepartmentIn(BaseModel):
    name: str


@router.post("/departments")
def create_department(payload: DepartmentIn, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))):
    name = payload.name.strip()
    if len(name) < 3:
        raise HTTPException(status_code=400, detail="Department name is too short")
    if db.query(Department).filter(Department.name.ilike(name)).first():
        raise HTTPException(status_code=400, detail="A department with this name already exists")
    d = Department(name=name)
    db.add(d)
    db.commit()
    db.refresh(d)
    log_activity(db, user, f"Created department {name}")
    return {"id": d.id, "name": d.name}


@router.patch("/departments/{department_id}")
def rename_department(department_id: int, payload: DepartmentIn, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))):
    d = db.query(Department).filter(Department.id == department_id).first()
    if not d:
        raise HTTPException(status_code=404, detail="Department not found")
    name = payload.name.strip()
    if len(name) < 3:
        raise HTTPException(status_code=400, detail="Department name is too short")
    if db.query(Department).filter(Department.name.ilike(name), Department.id != d.id).first():
        raise HTTPException(status_code=400, detail="A department with this name already exists")
    old, d.name = d.name, name
    db.commit()
    log_activity(db, user, f"Renamed department {old} to {name}")
    return {"id": d.id, "name": d.name}


@router.get("/structure")
def academic_structure(db: Session = Depends(get_db), user: User = Depends(require_role("admin"))):
    """Departments with their batches, semesters in use, and faculty and subject counts."""
    out = []
    for d in db.query(Department).order_by(Department.name).all():
        students = db.query(Student.batch, Student.semester, Student.status).filter(Student.department_id == d.id).all()
        batches: dict[str, dict] = {}
        for batch, semester, status in students:
            b = batches.setdefault(batch or "—", {"batch": batch or "—", "students": 0, "active": 0, "semesters": set()})
            b["students"] += 1
            if status == "active":
                b["active"] += 1
                if semester:
                    b["semesters"].add(semester)
        out.append({
            "id": d.id, "name": d.name,
            "faculty": db.query(Faculty).filter(Faculty.department_id == d.id).count(),
            "subjects": db.query(Subject).filter(Subject.department_id == d.id).count(),
            "subjects_without_faculty": db.query(Subject).filter(Subject.department_id == d.id, Subject.faculty_id.is_(None)).count(),
            "students": len(students),
            "batches": [{**b, "semesters": sorted(b["semesters"])} for b in sorted(batches.values(), key=lambda x: x["batch"], reverse=True)],
        })
    return out


class SubjectIn(BaseModel):
    code: Optional[str] = None
    name: Optional[str] = None
    semester: Optional[int] = None
    department_id: Optional[int] = None
    faculty_id: Optional[int] = None      # 0 clears the allocation
    scheme: Optional[str] = None


def _subject_out(s: Subject) -> dict:
    return {"id": s.id, "code": s.code, "name": s.name, "semester": s.semester, "scheme": s.scheme, "department_id": s.department_id,
            "faculty_id": s.faculty_id, "faculty_name": s.faculty.name if s.faculty else None, "name_verified": bool(s.name_verified)}


@router.get("/subjects/allocation")
def subject_allocation(db: Session = Depends(get_db), user: User = Depends(require_role("admin"))):
    return {
        "subjects": [_subject_out(s) for s in db.query(Subject).order_by(Subject.semester, Subject.code).all()],
        "faculty": [{"id": f.id, "name": f.name, "department_id": f.department_id} for f in db.query(Faculty).order_by(Faculty.name).all()],
    }


@router.post("/subjects")
def create_subject(payload: SubjectIn, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))):
    code, name = (payload.code or "").strip().upper(), (payload.name or "").strip()
    if not re.match(r"^[A-Z0-9]{4,12}$", code):
        raise HTTPException(status_code=400, detail="Subject code should be 4 to 12 letters and digits, e.g. BIS701")
    if len(name) < 3:
        raise HTTPException(status_code=400, detail="Subject name is required")
    if db.query(Subject).filter(Subject.code == code).first():
        raise HTTPException(status_code=400, detail="A subject with this code already exists")
    if payload.semester is not None and not 1 <= payload.semester <= 8:
        raise HTTPException(status_code=400, detail="Semester must be between 1 and 8")
    dept = db.query(Department).filter(Department.id == payload.department_id).first() if payload.department_id else db.query(Department).first()
    if not dept:
        raise HTTPException(status_code=400, detail="Department not found")
    if payload.faculty_id and not db.query(Faculty).filter(Faculty.id == payload.faculty_id).first():
        raise HTTPException(status_code=400, detail="Faculty member not found")
    s = Subject(code=code, name=name, semester=payload.semester, department_id=dept.id, faculty_id=payload.faculty_id or None,
                scheme=(payload.scheme or "").strip() or None, name_verified=True)
    db.add(s)
    db.commit()
    db.refresh(s)
    log_activity(db, user, f"Added subject {code} {name}")
    return _subject_out(s)


@router.patch("/subjects/{subject_id}")
def edit_subject(subject_id: int, payload: SubjectIn, db: Session = Depends(get_db), user: User = Depends(require_role("admin"))):
    s = db.query(Subject).filter(Subject.id == subject_id).first()
    if not s:
        raise HTTPException(status_code=404, detail="Subject not found")
    sent = payload.dict(exclude_unset=True)
    changes = []
    if "name" in sent and (sent["name"] or "").strip() and sent["name"].strip() != s.name:
        s.name, s.name_verified = sent["name"].strip(), True
        changes.append("renamed")
    if "semester" in sent and sent["semester"] != s.semester:
        if sent["semester"] is not None and not 1 <= sent["semester"] <= 8:
            raise HTTPException(status_code=400, detail="Semester must be between 1 and 8")
        s.semester = sent["semester"]
        changes.append(f"semester {sent['semester']}")
    if "faculty_id" in sent:
        new_id = sent["faculty_id"] or None
        if new_id != s.faculty_id:
            faculty = db.query(Faculty).filter(Faculty.id == new_id).first() if new_id else None
            if new_id and not faculty:
                raise HTTPException(status_code=400, detail="Faculty member not found")
            s.faculty_id = new_id
            changes.append(f"allocated to {faculty.name}" if faculty else "allocation removed")
    db.commit()
    db.refresh(s)
    if changes:
        log_activity(db, user, f"Subject {s.code}: {', '.join(changes)}")
    return _subject_out(s)


# ------------------------------------------------------------------ placement and internship records
# Open to the admin (everything) and to an HOD (their own department's students).

def _scope_student(db: Session, user: User, usn: str) -> Student:
    s = db.query(Student).filter(Student.usn == (usn or "").strip().upper()).first()
    if not s or (user.role == "hod" and s.department_id != user.department_id):
        raise HTTPException(status_code=404, detail="No student with that USN" + (" in your department" if user.role == "hod" else ""))
    return s


def _company(db: Session, name: str) -> Company:
    display, key = canonical_company(name or "")
    if not key:
        raise HTTPException(status_code=400, detail="Company name is required")
    c = db.query(Company).filter(Company.name_key == key).first()
    if not c:
        c = Company(name=display, name_key=key)
        db.add(c)
        db.flush()
    return c


def _positive(value, label: str) -> Decimal:
    try:
        number = Decimal(str(value))
    except (InvalidOperation, TypeError):
        raise HTTPException(status_code=400, detail=f"{label} must be a number")
    if number <= 0:
        raise HTTPException(status_code=400, detail=f"{label} must be greater than zero")
    return number


def _date(value: Optional[str], label: str):
    if not value:
        return None
    try:
        return datetime.strptime(value.strip(), "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=400, detail=f"{label} must be a date (YYYY-MM-DD)")


def _placement_out(p: Placement) -> dict:
    return {"id": p.id, "usn": p.student.usn, "name": p.student.name, "batch": p.student.batch,
            "company": p.company_ref.name if p.company_ref else p.company, "role": p.role,
            "package_lpa": float(p.package_lpa) if p.package_lpa is not None else None, "ctc_text": p.ctc,
            "offer_type": p.offer_type, "status": p.status, "offer_date": str(p.offer_date) if p.offer_date else None,
            "drive_type": p.drive.drive_type if p.drive else None, "verified": bool(p.verified), "placed": bool(p.placed)}


def _internship_out(i: Internship) -> dict:
    return {"id": i.id, "usn": i.student.usn, "name": i.student.name, "batch": i.student.batch,
            "company": i.company_ref.name if i.company_ref else i.company, "role": i.role,
            "stipend_per_month": float(i.stipend_per_month) if i.stipend_per_month is not None else None,
            "start_date": str(i.start_date) if i.start_date else None, "end_date": str(i.end_date) if i.end_date else None,
            "duration": i.duration_text, "status": i.status, "mode": i.mode, "verified": bool(i.verified)}


def _in_scope(query, user: User):
    return query.filter(Student.department_id == user.department_id) if user.role == "hod" else query


def announce_placements(db: Session, user: User, company: Company, batch: str, students: list[Student], package) -> None:
    """The placement wall: one notice per company and batch when offers are recorded."""
    if not students:
        return
    from app.core.placement_stats import batch_label
    from app.routers.campus import publish
    count = len(students)
    package_text = f" ({float(package):g} LPA)" if package is not None else ""
    names = ", ".join(s.name.title() for s in students[:12]) + (" and others" if count > 12 else "")
    publish(db, Announcement(
        author_user_id=user.id, author_name="Placement Cell", author_role=user.role, kind="placement", audience="everyone",
        department_id=students[0].department_id, batch=None, semester=None,
        title=f"{count} student{'s' if count != 1 else ''} placed at {company.name}{package_text}",
        body=f"Congratulations to {names} of the {batch_label(batch)} batch on {'their offers' if count != 1 else 'the offer'} from **{company.name}**{package_text}.",
    ))


@router.get("/placement-records")
def placement_records(batch: Optional[str] = None, db: Session = Depends(get_db), user: User = Depends(require_role("admin", "hod"))):
    pq = _in_scope(db.query(Placement).join(Student, Student.id == Placement.student_id).filter(Placement.placed == True), user)  # noqa: E712
    iq = _in_scope(db.query(Internship).join(Student, Student.id == Internship.student_id), user)
    if batch:
        pq, iq = pq.filter(Student.batch == batch), iq.filter(Student.batch == batch)
    batches = sorted({b for (b,) in _in_scope(db.query(Student.batch), user).distinct() if b}, reverse=True)
    return {
        "placements": [_placement_out(p) for p in pq.order_by(Student.batch.desc(), Student.usn).all()],
        "internships": [_internship_out(i) for i in iq.order_by(Student.batch.desc(), Student.usn).all()],
        "companies": [c.name for c in db.query(Company).order_by(Company.name).all()],
        "batches": batches, "offer_types": OFFER_TYPES, "statuses": OFFER_STATUSES,
    }


class PlacementIn(BaseModel):
    usn: Optional[str] = None
    company: Optional[str] = None
    role: Optional[str] = None
    package_lpa: Optional[float] = None
    offer_type: Optional[str] = None
    status: Optional[str] = None
    offer_date: Optional[str] = None
    drive_type: Optional[str] = None


def _apply_placement(db: Session, p: Placement, payload: PlacementIn, student: Student, creating: bool) -> None:
    sent = payload.dict(exclude_unset=True)
    if creating or "company" in sent:
        company = _company(db, payload.company or "")
        p.company_id, p.company = company.id, company.name
    if creating or "package_lpa" in sent:
        p.package_lpa = _positive(payload.package_lpa, "Package")
        p.ctc = f"{p.package_lpa.normalize():f} LPA"
    if creating or "offer_type" in sent:
        if payload.offer_type not in OFFER_TYPES:
            raise HTTPException(status_code=400, detail=f"Offer type must be one of: {', '.join(OFFER_TYPES)}")
        p.offer_type = payload.offer_type
    if creating or "status" in sent:
        status = payload.status or "offered"
        if status not in OFFER_STATUSES:
            raise HTTPException(status_code=400, detail=f"Status must be one of: {', '.join(OFFER_STATUSES)}")
        p.status = status
    if "role" in sent:
        p.role = (payload.role or "").strip() or None
    if "offer_date" in sent:
        p.offer_date = _date(payload.offer_date, "Offer date")
    if creating or "drive_type" in sent:
        drive_type = (payload.drive_type or "").strip() or "Not specified"
        drive = db.query(PlacementDrive).filter(PlacementDrive.company_id == p.company_id, PlacementDrive.batch == student.batch,
                                                PlacementDrive.drive_type == drive_type, PlacementDrive.drive_date.is_(None)).first()
        if not drive:
            drive = PlacementDrive(company_id=p.company_id, batch=student.batch, drive_type=drive_type)
            db.add(drive)
            db.flush()
        p.drive_id = drive.id
    p.placed, p.verified, p.usn = True, True, student.usn
    clash = db.query(Placement).filter(Placement.student_id == student.id, Placement.company_id == p.company_id,
                                       Placement.offer_type == p.offer_type, Placement.id != (p.id or 0)).first()
    if clash:
        raise HTTPException(status_code=400, detail="This student already has that kind of offer from this company")


@router.post("/placement-records/placements")
def add_placement(payload: PlacementIn, db: Session = Depends(get_db), user: User = Depends(require_role("admin", "hod"))):
    student = _scope_student(db, user, payload.usn or "")
    p = Placement(student_id=student.id, created_by=user.id, source_note=f"Entered by {user.email}")
    _apply_placement(db, p, payload, student, creating=True)
    db.add(p)
    db.commit()
    db.refresh(p)
    log_activity(db, user, f"Added placement: {student.usn} at {p.company} ({p.ctc})")
    announce_placements(db, user, p.company_ref, student.batch, [student], p.package_lpa)
    return _placement_out(p)


def _placement_or_404(db: Session, user: User, placement_id: int) -> Placement:
    p = _in_scope(db.query(Placement).join(Student, Student.id == Placement.student_id), user).filter(Placement.id == placement_id).first()
    if not p:
        raise HTTPException(status_code=404, detail="Placement record not found")
    return p


@router.patch("/placement-records/placements/{placement_id}")
def edit_placement(placement_id: int, payload: PlacementIn, db: Session = Depends(get_db), user: User = Depends(require_role("admin", "hod"))):
    p = _placement_or_404(db, user, placement_id)
    _apply_placement(db, p, payload, p.student, creating=False)
    db.commit()
    db.refresh(p)
    log_activity(db, user, f"Edited placement: {p.student.usn} at {p.company} ({p.ctc}, {p.status})")
    return _placement_out(p)


@router.delete("/placement-records/placements/{placement_id}")
def delete_placement(placement_id: int, db: Session = Depends(get_db), user: User = Depends(require_role("admin", "hod"))):
    p = _placement_or_404(db, user, placement_id)
    summary = f"{p.student.usn} at {p.company} ({p.ctc})"
    db.delete(p)
    db.commit()
    log_activity(db, user, f"Deleted placement: {summary}")
    return {"message": "Deleted"}


class InternshipIn(BaseModel):
    usn: Optional[str] = None
    company: Optional[str] = None
    role: Optional[str] = None
    stipend_per_month: Optional[float] = None
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    status: Optional[str] = None
    mode: Optional[str] = None


def _apply_internship(db: Session, i: Internship, payload: InternshipIn, student: Student, creating: bool) -> None:
    sent = payload.dict(exclude_unset=True)
    if creating or "company" in sent:
        company = _company(db, payload.company or "")
        i.company_id, i.company = company.id, company.name
    if "stipend_per_month" in sent:
        i.stipend_per_month = _positive(payload.stipend_per_month, "Stipend") if payload.stipend_per_month not in (None, 0) else None
    if "start_date" in sent:
        i.start_date = _date(payload.start_date, "Start date")
    if "end_date" in sent:
        i.end_date = _date(payload.end_date, "End date")
    if i.start_date and i.end_date and i.end_date < i.start_date:
        raise HTTPException(status_code=400, detail="End date is before the start date")
    if "status" in sent:
        if payload.status and payload.status not in ("ongoing", "completed"):
            raise HTTPException(status_code=400, detail="Status must be ongoing or completed")
        i.status = payload.status or None
    if "role" in sent:
        i.role = (payload.role or "").strip() or None
    if "mode" in sent:
        i.mode = (payload.mode or "").strip() or None
    i.verified, i.usn = True, student.usn
    same = db.query(Internship).filter(Internship.student_id == student.id, Internship.company_id == i.company_id, Internship.id != (i.id or 0))
    same = same.filter(Internship.start_date == i.start_date) if i.start_date else same.filter(Internship.start_date.is_(None))
    if same.first():
        raise HTTPException(status_code=400, detail="This internship is already recorded for the student")


@router.post("/placement-records/internships")
def add_internship(payload: InternshipIn, db: Session = Depends(get_db), user: User = Depends(require_role("admin", "hod"))):
    student = _scope_student(db, user, payload.usn or "")
    i = Internship(student_id=student.id, source_note=f"Entered by {user.email}", company="")
    _apply_internship(db, i, payload, student, creating=True)
    db.add(i)
    db.commit()
    db.refresh(i)
    log_activity(db, user, f"Added internship: {student.usn} at {i.company}")
    return _internship_out(i)


def _internship_or_404(db: Session, user: User, internship_id: int) -> Internship:
    i = _in_scope(db.query(Internship).join(Student, Student.id == Internship.student_id), user).filter(Internship.id == internship_id).first()
    if not i:
        raise HTTPException(status_code=404, detail="Internship record not found")
    return i


@router.patch("/placement-records/internships/{internship_id}")
def edit_internship(internship_id: int, payload: InternshipIn, db: Session = Depends(get_db), user: User = Depends(require_role("admin", "hod"))):
    i = _internship_or_404(db, user, internship_id)
    _apply_internship(db, i, payload, i.student, creating=False)
    db.commit()
    db.refresh(i)
    log_activity(db, user, f"Edited internship: {i.student.usn} at {i.company}")
    return _internship_out(i)


@router.delete("/placement-records/internships/{internship_id}")
def delete_internship(internship_id: int, db: Session = Depends(get_db), user: User = Depends(require_role("admin", "hod"))):
    i = _internship_or_404(db, user, internship_id)
    summary = f"{i.student.usn} at {i.company}"
    db.delete(i)
    db.commit()
    log_activity(db, user, f"Deleted internship: {summary}")
    return {"message": "Deleted"}


@router.post("/placement-records/import")
async def import_placement_records(
    placements: Optional[UploadFile] = File(None), internships: Optional[UploadFile] = File(None), apply: bool = Form(False),
    db: Session = Depends(get_db), user: User = Depends(require_role("admin")),
):
    """CSV import with the same strict student matching as the command-line
    importer. Without `apply` it only reports what each row would do."""
    from scripts.import_placements import run

    if not placements and not internships:
        raise HTTPException(status_code=400, detail="Choose a placements file, an internships file, or both")
    paths = {}
    try:
        for label, upload in (("placements", placements), ("internships", internships)):
            if upload is None:
                paths[label] = ""
                continue
            if not (upload.filename or "").lower().endswith(".csv"):
                raise HTTPException(status_code=400, detail="Upload .csv files (columns as in the template)")
            with tempfile.NamedTemporaryFile(suffix=".csv", delete=False) as tmp:
                tmp.write(await upload.read())
                paths[label] = tmp.name
        before = {pid for (pid,) in db.query(Placement.id).all()}
        result = run(paths["placements"], paths["internships"], None, apply, db=db)
    except SystemExit as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    except KeyError as exc:
        raise HTTPException(status_code=400, detail=f"A required column is missing: {exc}")
    finally:
        for path in paths.values():
            if path and os.path.exists(path):
                os.remove(path)

    def rows(kind: str) -> list[dict]:
        out = []
        for e in result[kind]:
            m = e["match"]
            out.append({
                "name": e["row"].get("name"), "company": canonical_company(e["row"].get("company") or "")[0],
                "amount": e["row"].get("package_lpa") if kind == "placements" else e["row"].get("stipend_per_month"),
                "decision": e["decision"], "matched_usn": m.student.usn if m.student else None, "matched_name": m.student.name if m.student else None,
                "how": m.method, "reason": "; ".join(e["errors"]) or m.reason or e["open_question"],
                "candidates": [{"usn": c.usn, "name": c.name, "score": c.score} for c in m.candidates[:5]],
            })
        return out

    if apply:
        log_activity(db, user, f"Imported placement records: {result['written']['placements']} placements, {result['written']['internships']} internships")
        new = db.query(Placement).filter(~Placement.id.in_(before)).all() if before else db.query(Placement).all()
        groups: dict[tuple, list[Placement]] = {}
        for p in new:
            if p.company_ref:
                groups.setdefault((p.company_id, p.student.batch), []).append(p)
        for (_, batch), offers in groups.items():
            announce_placements(db, user, offers[0].company_ref, batch, [o.student for o in offers], max((o.package_lpa for o in offers if o.package_lpa is not None), default=None))
    return {"applied": bool(apply), "written": result["written"], "placements": rows("placements"), "internships": rows("internships")}
