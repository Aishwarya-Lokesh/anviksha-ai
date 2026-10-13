"""The report engine: reading a format, and writing a report into it without
changing the format."""
import io

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.core import project_docx as pd
from app.core import project_fields as pf
from tests import project_samples as S


def _texts(data: bytes) -> list[str]:
    doc = Document(io.BytesIO(data))
    return [pd.text_of(el).strip() for el in doc.element.body.iter(pd.P) if pd.text_of(el).strip()]


def _built(template=None, content=None, details=None, **kw):
    template = template or S.report_template()
    profile = pd.analyse(template)
    return profile, pd.build(template, profile["field_map"], details or S.DETAILS, content or S.CONTENT, {1: S.png(), 2: S.png(600, 900)}, **kw)


def test_analysis_reads_the_format():
    a = pd.analyse(S.report_template())
    assert a["page"] == {"width_in": 8.27, "height_in": 11.69, "margins_in": {"top": 1.0, "bottom": 1.0, "left": 1.25, "right": 1.0}}
    assert [(s["format"], s["start"]) for s in a["sections"]] == [("lowerRoman", 1), ("decimal", 1)]
    assert a["parts"] == ["Cover page", "Certificate", "Declaration", "Acknowledgement", "Abstract", "Table of contents", "List of figures", "List of tables", "Chapters", "References"]
    assert a["chapter_titles"] == ["Introduction", "Literature Survey"]
    assert a["styles"]["body text"] == {"font": "Times New Roman", "size": 12.0, "bold": False, "italic": False, "align": "justified", "line_spacing": 1.5}
    assert (a["styles"]["chapter heading"]["size"], a["styles"]["chapter heading"]["align"], a["styles"]["section heading"]["size"]) == (16.0, "centre", 14.0)
    assert a["captions"] == {"figure": "Fig 4.1:", "table": "Table 4.1:", "table_caption_position": "above the table"}
    assert any("[page number]" in f for f in a["footers"]) and a["headers"] == ["TITLE OF THE PROJECT"]
    fields = {m["match"]: m["field"] for m in a["field_map"]}
    assert fields == {"TITLE OF THE PROJECT": "title", "NAME OF THE STUDENT (USN)": "student_name_usn", "<<Name of Guide>>": "guide_name",
                      "<<Guide Designation>>": "guide_designation", "<<Name of HOD>>": "hod_name", "<<Name of Principal>>": "principal_name",
                      "[Student Name]": "student_name", "[USN]": "student_usn", "20XX-20XX": "academic_year"}
    assert a["warnings"] == []
    with pytest.raises(ValueError):
        pd.analyse(b"not a document")


def test_fields_are_filled_everywhere_including_header_footer_and_tables():
    _, out = _built()
    doc = Document(io.BytesIO(out["docx"]))
    texts = _texts(out["docx"])
    assert "“SMART ATTENDANCE SYSTEM USING FACE RECOGNITION”" in texts                       # upper case, like the format's wording
    assert texts[texts.index("Submitted by") + 1: texts.index("Under the guidance of")] == ["ASHA RAO (4GW22IS001)", "BINA KUMAR (4GW22IS002)", "CHITRA N (4GW22IS003)"]
    certificate = next(t for t in texts if t.startswith("Certified that"))
    assert "carried out by ASHA RAO (4GW22IS001), BINA KUMAR (4GW22IS002) and CHITRA N (4GW22IS003) in partial" in certificate and "2025-2026" in certificate
    team = doc.tables[1]
    assert [[c.text for c in row.cells] for row in team.rows] == [["Name", "USN"], ["Asha Rao", "4GW22IS001"], ["Bina Kumar", "4GW22IS002"], ["Chitra N", "4GW22IS003"]]
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["Dr. Meera Nair", "Dr. Latha Rao", "Dr. Shiva Kumar"]
    assert doc.sections[1].header.paragraphs[0].text == "SMART ATTENDANCE SYSTEM USING FACE RECOGNITION"
    assert "2025-2026" in doc.sections[1].footer.paragraphs[0].text
    assert out["unfilled"] == [] and pd.leftovers(out["docx"], pd.analyse(S.report_template())["field_map"]) == []


def test_missing_details_are_reported_not_invented():
    profile, out = _built(details={"title": "Only A Title", "students": []})
    assert set(out["unfilled"]) == {"Student name with USN (one per student)", "Guide's name", "Guide's designation", "HOD's name", "Principal's name",
                                    "Student name (one per student)", "Student USN (one per student)", "Academic year"}
    assert "NAME OF THE STUDENT (USN)" in _texts(out["docx"])                              # left as the format has it, never made up
    assert "<<Name of Guide>>" in pd.leftovers(out["docx"], profile["field_map"])


def test_content_is_written_in_the_formats_own_paragraphs():
    template = S.report_template()
    _, out = _built(template)
    doc = Document(io.BytesIO(out["docx"]))
    texts = _texts(out["docx"])
    assert not any("sample" in t.lower() for t in texts)                                    # none of the format's sample content survives
    start = texts.index("CHAPTER 1")
    assert texts[start: start + 4] == ["CHAPTER 1", "INTRODUCTION", "1.1 Overview", "Attendance is recorded in every class. The manual roll call takes several minutes of each session."]
    for want in ("1.2 Problem Statement", "1.2.1 Scope", "CHAPTER 2", "SYSTEM DESIGN", "2.1 Architecture", "Fig 2.1: System architecture", "Table 2.1: Modules and their roles",
                 "Fig 2.2: Data flow", "CHAPTER 3", "CONCLUSION"):
        assert want in texts, want
    assert "LITERATURE SURVEY" not in texts
    assert texts[texts.index("REFERENCES") + 1:] == ["[1] P. Viola and M. Jones, “Rapid object detection using a boosted cascade of simple features,” CVPR, 2001.",
                                                    "[2] F. Schroff, D. Kalenichenko and J. Philbin, “FaceNet: A unified embedding for face recognition and clustering,” CVPR, 2015."]

    by_text = {pd.text_of(el).strip(): el for el in doc.element.body.iter(pd.P)}
    look = lambda t: pd._resolved(doc, by_text[t])  # noqa: E731
    assert look("The system has a capture module, a recognition module and a records module.") == {"font": "Times New Roman", "size": 12.0, "bold": False, "italic": False, "align": "justified", "line_spacing": 1.5}
    assert (look("SYSTEM DESIGN")["size"], look("SYSTEM DESIGN")["bold"], look("SYSTEM DESIGN")["align"]) == (16.0, True, "centre")
    assert (look("2.1 Architecture")["size"], look("2.1 Architecture")["bold"]) == (14.0, True)
    assert (look("Fig 2.1: System architecture")["size"], look("Fig 2.1: System architecture")["align"]) == (11.0, "centre")
    bold_run = next(r for r in by_text["Attendance is recorded in every class. The manual roll call takes several minutes of each session."].iter(pd.R) if pd.text_of(r) == "manual roll call")
    assert bold_run.find(pd.RPR).find(qn("w:b")) is not None
    assert len(doc.inline_shapes) == 2
    page_height = doc.sections[0].page_height
    assert all(sh.height <= page_height * 0.43 and sh.width <= Inches_text_width(doc) for sh in doc.inline_shapes)   # a tall image is scaled to fit the page
    module_table = doc.tables[-1]
    assert [[c.text for c in row.cells] for row in module_table.rows][1] == ["Capture", "Reads frames from the camera"]


def Inches_text_width(doc):
    s = doc.sections[-1]
    return s.page_width - s.left_margin - s.right_margin


def test_page_setup_sections_and_numbering_are_untouched():
    template = S.report_template()
    before, after = Document(io.BytesIO(template)), Document(io.BytesIO(_built(template)[1]["docx"]))
    assert len(after.sections) == len(before.sections) == 2
    for a, b in zip(before.sections, after.sections):
        assert (a.page_width, a.page_height, a.top_margin, a.bottom_margin, a.left_margin, a.right_margin) == (b.page_width, b.page_height, b.top_margin, b.bottom_margin, b.left_margin, b.right_margin)
    assert pd._numbering(after) == pd._numbering(before) == [{"section": 1, "format": "lowerRoman", "start": 1}, {"section": 2, "format": "decimal", "start": 1}]
    # the chapters are still in the second section: the section break sits before CHAPTER 1
    order = [("break" if pd._has_sect(el) else pd.text_of(el).strip()) for el in after.element.body.iter(pd.P)]
    assert order.index("break") < order.index("CHAPTER 1") and order.index("LIST OF TABLES") < order.index("break")
    assert any("PAGE" in (i.text or "") for i in after.sections[1].footer._element.iter(qn("w:instrText")))


def test_contents_lists_what_was_written_with_page_numbers_when_known():
    _, first = _built(markers=True)
    assert [e["text"] for e in first["toc"]] == ["Acknowledgement", "Abstract", "Chapter 1: Introduction", "1.1 Overview", "1.2 Problem Statement", "1.2.1 Scope",
                                                 "Chapter 2: System Design", "2.1 Architecture", "Chapter 3: Conclusion", "References"]
    assert [(f["label"], f["caption"]) for f in first["figures"]] == [("Fig 2.1", "System architecture"), ("Fig 2.2", "Data flow")]
    assert [t["label"] for t in first["tables"]] == ["Table 2.1"]
    tokens = {a["token"] for a in first["anchors"]} | {m["token"] for m in first["section_marks"]}
    assert len(first["section_marks"]) == 2 and all(any(tok in t for t in _texts(first["docx"])) for tok in tokens)

    pages = {"front:abstract": "ii", "chapter:1": "1", "heading:1.1": "1", "chapter:2": "3", "figure:2.1": "3", "table:2.1": "4", "front:references": "6"}
    profile, final = _built(pages=pages)
    texts = _texts(final["docx"])
    contents = texts[texts.index("TABLE OF CONTENTS") + 1: texts.index("LIST OF FIGURES")]
    assert contents[:4] == ["Acknowledgement", "Abstractii", "Chapter 1: Introduction1", "1.1 Overview1"] and "Chapter 2: System Design3" in contents and contents[-1] == "References6"
    assert texts[texts.index("LIST OF FIGURES") + 1: texts.index("LIST OF TABLES")] == ["Fig 2.1 System architecture3", "Fig 2.2 Data flow"]
    assert texts[texts.index("LIST OF TABLES") + 1] == "Table 2.1 Modules and their roles4"
    assert pd.leftovers(final["docx"], profile["field_map"]) == []                        # no page-finder tokens in the final file
    toc_line = next(el for el in Document(io.BytesIO(final["docx"])).element.body.iter(pd.P) if pd.text_of(el).startswith("Chapter 2: System Design"))
    assert toc_line.find(pd.PPR).find(qn("w:tabs")) is not None and len(list(toc_line.iter(qn("w:tab")))) >= 2   # leader tab kept, title and number separated by a tab


def test_other_layouts_of_format_are_handled():
    # chapters marked only by the Heading 1 style
    template = S.report_template(heading1_style=True)
    a = pd.analyse(template)
    assert a["chapter_titles"] == ["Introduction"] and a["warnings"] == []
    texts = _texts(pd.build(template, a["field_map"], S.DETAILS, S.CONTENT, {1: S.png(), 2: S.png()})["docx"])
    assert {"1. INTRODUCTION", "2. SYSTEM DESIGN", "3. CONCLUSION", "2.1 Architecture"} <= set(texts)
    doc = Document(io.BytesIO(pd.build(template, a["field_map"], S.DETAILS, S.CONTENT, {})["docx"]))
    heading = next(p for p in doc.paragraphs if p.text == "2. SYSTEM DESIGN")
    assert heading.style.name == "Heading 1"

    # contents page laid out as a table
    template = S.report_template(contents_as_table=True)
    a = pd.analyse(template)
    doc = Document(io.BytesIO(pd.build(template, a["field_map"], S.DETAILS, S.CONTENT, {}, pages={"chapter:2": "3"})["docx"]))
    toc = next(t for t in doc.tables if t.cell(0, 0).text == "Chapter No.")
    rows = [[c.text for c in r.cells] for r in toc.rows]
    assert rows[0] == ["Chapter No.", "Title", "Page No."] and ["2", "System Design", "3"] in rows and ["1.1", "Overview", ""] in rows

    # a missing image is dropped and reported, never replaced by a placeholder
    out = pd.build(S.report_template(), a["field_map"], S.DETAILS, S.CONTENT, {})
    assert out["dropped"] == ["System architecture", "Data flow"] and out["figures"] == []


def test_field_map_can_be_corrected_by_a_person():
    assert pf.clean_map([{"match": " Sample Title ", "field": "title"}, {"match": "x", "field": "title"}, {"match": "Guide", "field": "nonsense"},
                         {"match": "Sample Title", "field": "title"}]) == [{"match": "Sample Title", "field": "title"}]
    assert pf.guess_map(["Project by <Student Name> under [Guide Name], year 2024-25"]) == [
        {"match": "<Student Name>", "field": "student_name"}, {"match": "[Guide Name]", "field": "guide_name"}, {"match": "2024-25", "field": "academic_year"}]
    assert pf.inline_value("student_name", "name", {"students": [{"name": "A", "usn": ""}, {"name": "B", "usn": ""}]}) == "A and B"
