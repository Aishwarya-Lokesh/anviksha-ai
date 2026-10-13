"""PowerPoint export: themes, stock photos, and the designed panel used when a
slide has no photo."""
import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.core import slide_images
from tests.test_missing_endpoints import campus  # noqa: F401  (shared fixture)

SLIDES = [
    {"title": "Introduction to Computer Networks", "content": "- Nodes and links\n- **Protocols** define the rules"},
    {"title": "OSI Model Layers", "content": "- Seven layers", "diagram_description": "A stack of seven labelled layers"},
    {"title": "TCP Handshake", "content": "- SYN\n- SYN-ACK\n- ACK", "process_steps": ["SYN", "SYN-ACK", "ACK"]},
]


def _png() -> bytes:
    from PIL import Image
    buffer = io.BytesIO()
    Image.new("RGB", (900, 900), (40, 90, 160)).save(buffer, format="PNG")
    return buffer.getvalue()


def _export(client, headers, **extra):
    r = client.post("/teacher/export-ppt", headers=headers, json={"topic": "Computer Networks", "slides": SLIDES, **extra})
    assert r.status_code == 200, r.text
    return Presentation(io.BytesIO(r.content))


def _pictures(prs) -> list[int]:
    return [sum(1 for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE) for slide in prs.slides]


def _texts(slide) -> str:
    return " | ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def test_export_without_a_photo_key_uses_designed_panels(client, campus, monkeypatch):  # noqa: F811
    monkeypatch.delenv("PEXELS_API_KEY", raising=False)
    monkeypatch.delenv("PIXABAY_API_KEY", raising=False)
    h = campus["h"]["t1"]
    info = client.get("/teacher/slide-themes", headers=h).json()
    assert info["images_available"] is False and [t["key"] for t in info["themes"]] == ["violet", "ocean", "forest", "sunset", "slate"]
    assert client.get("/teacher/slide-image?q=OSI model", headers=h).json() == {"available": False, "reason": "not_configured"}

    prs = _export(client, h)
    assert len(prs.slides) == 4 and _pictures(prs) == [0, 0, 0, 0]
    assert "PICTURE THIS" in _texts(prs.slides[2]) and "A stack of seven labelled layers" in _texts(prs.slides[2])
    assert "SUGGESTED VISUAL" not in " ".join(_texts(s) for s in prs.slides)        # no placeholder box
    assert "1. SYN" in _texts(prs.slides[3])                                        # process flow untouched


def test_theme_changes_the_deck_colour(client, campus, monkeypatch):  # noqa: F811
    monkeypatch.delenv("PEXELS_API_KEY", raising=False)
    h = campus["h"]["t1"]
    colour = lambda prs: str(prs.slides[0].shapes[0].fill.fore_color.rgb)  # noqa: E731
    assert colour(_export(client, h)) == "5B3FD9"
    assert colour(_export(client, h, theme="ocean")) == "0E7490"
    assert colour(_export(client, h, theme="no-such-theme")) == "5B3FD9"


def test_photos_are_placed_with_credit_and_only_from_the_photo_host(client, campus, monkeypatch):  # noqa: F811
    monkeypatch.setenv("PEXELS_API_KEY", "test-key")
    slide_images._cache.clear()
    asked = []

    def fake_search(query, index=0):
        asked.append(query)
        return {"url": "https://images.pexels.com/photos/1/a.jpeg", "thumb": "https://images.pexels.com/photos/1/t.jpeg", "credit": "Asha Rao / Pexels", "credit_url": "https://www.pexels.com/@asha", "alt": query}

    monkeypatch.setattr(slide_images, "search", fake_search)
    monkeypatch.setattr(slide_images, "fetch", lambda url: _png() if slide_images.allowed_url(url) else None)
    h = campus["h"]["t1"]

    found = client.get("/teacher/slide-image?q=Introduction to Computer Networks", headers=h).json()
    assert found["available"] is True and found["credit"] == "Asha Rao / Pexels" and asked[-1] == "Computer Networks"

    prs = _export(client, h)
    assert _pictures(prs) == [0, 1, 1, 0]                                           # not on the title slide or the process slide
    assert "Photo: Asha Rao / Pexels" in _texts(prs.slides[1])
    picture = next(sh for sh in prs.slides[1].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert picture.crop_top > 0 and picture.height <= 4.0 * 914400                   # a square photo is trimmed to the panel, not squashed

    assert _pictures(_export(client, h, include_images=False)) == [0, 0, 0, 0]
    assert client.get("/teacher/slide-image?q=x", headers=campus["h"]["s1"]).status_code == 403


def test_either_photo_service_can_be_used(monkeypatch):
    calls = []

    class Reply:
        def __init__(self, body):
            self.body = body

        def raise_for_status(self):
            return None

        def json(self):
            return self.body

    def fake_get(url, params=None, headers=None, timeout=None, **_):
        calls.append((url, params, headers))
        if "pixabay" in url:
            return Reply({"hits": [{"largeImageURL": "https://pixabay.com/get/big.jpg", "webformatURL": "https://pixabay.com/get/small.jpg", "user": "Meera", "pageURL": "https://pixabay.com/photos/1", "tags": "network, cable"}]})
        return Reply({"photos": [{"src": {"large": "https://images.pexels.com/photos/1/large.jpeg", "medium": "https://images.pexels.com/photos/1/medium.jpeg"}, "photographer": "Asha Rao", "photographer_url": "https://www.pexels.com/@asha", "alt": "cables"}]})

    monkeypatch.setattr(slide_images.httpx, "get", fake_get)
    monkeypatch.delenv("PEXELS_API_KEY", raising=False)
    monkeypatch.delenv("PIXABAY_API_KEY", raising=False)
    slide_images._cache.clear()
    assert slide_images.configured() is False and slide_images.search("network cable") is None and calls == []

    monkeypatch.setenv("PIXABAY_API_KEY", "pix-key")
    found = slide_images.search("network cable")
    assert (found["url"], found["thumb"], found["credit"]) == ("https://pixabay.com/get/big.jpg", "https://pixabay.com/get/small.jpg", "Meera / Pixabay")
    assert calls[-1][1]["key"] == "pix-key" and calls[-1][1]["safesearch"] == "true" and slide_images.allowed_url(found["url"])
    assert slide_images.search("network cable", 5) is None and len(calls) == 1          # answered from memory, no second request

    monkeypatch.delenv("PIXABAY_API_KEY")
    monkeypatch.setenv("PEXELS_API_KEY", "pex-key")
    found = slide_images.search("network cable")
    assert found["credit"] == "Asha Rao / Pexels" and calls[-1][2] == {"Authorization": "pex-key"}   # the other service's answer is not reused


def test_a_chosen_image_must_come_from_the_photo_host(monkeypatch):
    assert slide_images.allowed_url("https://images.pexels.com/photos/1/a.jpeg")
    assert slide_images.allowed_url("https://pixabay.com/get/abc.jpg")
    for bad in ("http://images.pexels.com/a.jpg", "https://images.pexels.com.evil.test/a.jpg", "https://pixabay.com.evil.test/a.jpg", "https://169.254.169.254/latest", "file:///etc/passwd", ""):
        assert not slide_images.allowed_url(bad)
    assert slide_images.fetch("https://example.com/a.jpg") is None
    assert slide_images.query_for({"title": "Introduction to the OSI Model"}) == "OSI Model"
    assert slide_images.query_for({"title": "Overview", "image_query": "fibre optic cable"}) == "fibre optic cable"
    assert slide_images.query_for({"title": "Summary"}, "Database Normalisation") == "Database Normalisation"
