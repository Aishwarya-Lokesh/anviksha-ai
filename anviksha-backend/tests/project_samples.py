"""Stand-in project formats and content for tests.

These are NOT the college's official formats. They are built here to look
and behave like a typical engineering-college report and presentation format
(cover, certificate, declaration, roman then arabic page numbers, "CHAPTER 1"
headings, "Fig 1.1" captions, a slide master with a header band), so the
engine can be tested against something realistic."""
import io

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SAMPLE_BODY = ("This is sample body text of the format. It shows the font, the size, the line spacing and the alignment that every "
               "paragraph of a chapter must follow. Replace it with the content of the project.")


def _field(paragraph, code: str):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
            run._r.append(node)
        elif text == code:
            node = OxmlElement("w:instrText")
            node.text = f" {code} "
            run._r.append(node)
        else:
            node = OxmlElement("w:t")
            node.text = text
            run._r.append(node)


def _page_numbers(section, fmt: str, start=None):
    node = section._sectPr.find(qn("w:pgNumType"))
    if node is None:
        node = OxmlElement("w:pgNumType")
        section._sectPr.append(node)
    node.set(qn("w:fmt"), fmt)
    if start is not None:
        node.set(qn("w:start"), str(start))


def _para(doc, text, size=12, bold=False, align=None, italic=False, spacing=1.5, after=6, first_indent=None, style=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.font.name, run.font.size, run.font.bold, run.font.italic = "Times New Roman", Pt(size), bold, italic
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    fmt = p.paragraph_format
    fmt.line_spacing, fmt.space_after = spacing, Pt(after)
    if align is not None:
        fmt.alignment = align
    if first_indent:
        fmt.first_line_indent = Inches(first_indent)
    return p


def _break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def report_template(heading1_style: bool = False, contents_as_table: bool = False) -> bytes:
    doc = Document()
    C, J = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Inches(1), Inches(1), Inches(1.25), Inches(1)
    s.different_first_page_header_footer = True
    _page_numbers(s, "lowerRoman", 1)
    foot = s.footer.paragraphs[0]
    foot.alignment = C
    _field(foot, "PAGE")

    # cover
    _para(doc, "VISVESVARAYA TECHNOLOGICAL UNIVERSITY", 14, True, C)
    _para(doc, "Jnana Sangama, Belagavi - 590018", 12, False, C)
    _para(doc, "A Project Report on", 12, False, C, italic=True)
    _para(doc, "“TITLE OF THE PROJECT”", 16, True, C)
    _para(doc, "Submitted in partial fulfilment of the requirements for the award of the degree of", 12, False, C)
    _para(doc, "BACHELOR OF ENGINEERING in INFORMATION SCIENCE AND ENGINEERING", 12, True, C)
    _para(doc, "Submitted by", 12, False, C)
    _para(doc, "NAME OF THE STUDENT (USN)", 12, True, C)
    _para(doc, "Under the guidance of", 12, False, C)
    _para(doc, "<<Name of Guide>>", 12, True, C)
    _para(doc, "<<Guide Designation>>", 12, False, C)
    _para(doc, "Department of Information Science and Engineering", 13, True, C)
    _para(doc, "GSSS Institute of Engineering and Technology for Women, Mysuru", 13, True, C)
    _para(doc, "20XX-20XX", 12, True, C)
    _break(doc)

    # certificate
    _para(doc, "CERTIFICATE", 16, True, C)
    _para(doc, "Certified that the project work entitled “TITLE OF THE PROJECT” is a bonafide work carried out by NAME OF THE STUDENT (USN) "
               "in partial fulfilment for the award of Bachelor of Engineering in Information Science and Engineering of the Visvesvaraya "
               "Technological University, Belagavi during the year 20XX-20XX.", 12, False, J)
    table = doc.add_table(rows=2, cols=3)
    for i, text in enumerate(("Signature of the Guide", "Signature of the HOD", "Signature of the Principal")):
        table.cell(0, i).text = text
    for i, text in enumerate(("<<Name of Guide>>", "<<Name of HOD>>", "<<Name of Principal>>")):
        table.cell(1, i).text = text
    _break(doc)

    # declaration, with one table row per student
    _para(doc, "DECLARATION", 16, True, C)
    _para(doc, "We, the students of final year Information Science and Engineering, hereby declare that the project entitled "
               "“TITLE OF THE PROJECT” has been carried out by us under the guidance of <<Name of Guide>>.", 12, False, J)
    team = doc.add_table(rows=2, cols=2)
    team.cell(0, 0).text, team.cell(0, 1).text = "Name", "USN"
    team.cell(1, 0).text, team.cell(1, 1).text = "[Student Name]", "[USN]"
    _break(doc)

    _para(doc, "ACKNOWLEDGEMENT", 16, True, C)
    _para(doc, "Sample acknowledgement text of the format, thanking the guide, the head of the department and the institution.", 12, False, J, first_indent=0.5)
    _break(doc)
    _para(doc, "ABSTRACT", 16, True, C)
    _para(doc, "Sample abstract text of the format. It is a single justified paragraph that summarises the problem, the approach and the outcome.", 12, False, J, first_indent=0.5)
    _break(doc)

    _para(doc, "TABLE OF CONTENTS", 16, True, C)
    if contents_as_table:
        toc = doc.add_table(rows=2, cols=3)
        for i, text in enumerate(("Chapter No.", "Title", "Page No.")):
            toc.cell(0, i).text = text
        for i, text in enumerate(("1", "Introduction", "1")):
            toc.cell(1, i).text = text
    else:
        for line in ("Abstract\tii", "Chapter 1: Introduction\t1", "1.1 Overview\t1"):
            p = _para(doc, line, 12, False, None, spacing=1.15)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.0), alignment=2, leader=1)
    _break(doc)
    _para(doc, "LIST OF FIGURES", 16, True, C)
    p = _para(doc, "Fig 1.1 Sample figure\t2", 12, False, None, spacing=1.15)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.0), alignment=2, leader=1)
    _break(doc)
    _para(doc, "LIST OF TABLES", 16, True, C)
    p = _para(doc, "Table 1.1 Sample table\t3", 12, False, None, spacing=1.15)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.0), alignment=2, leader=1)

    # chapters: a new section, arabic numbers from 1, header with the title
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.different_first_page_header_footer = False
    body.header.is_linked_to_previous = False
    body.footer.is_linked_to_previous = False
    _page_numbers(body, "decimal", 1)
    head = body.header.paragraphs[0]
    hr = head.add_run("TITLE OF THE PROJECT")
    hr.font.size, hr.font.italic, hr.font.name = Pt(10), True, "Times New Roman"
    f = body.footer.paragraphs[0]
    fr = f.add_run("Dept. of ISE, GSSSIETW\t20XX-20XX\tPage ")
    fr.font.size, fr.font.name = Pt(10), "Times New Roman"
    _field(f, "PAGE")

    if heading1_style:
        h = doc.add_heading("1. INTRODUCTION", level=1)
        h.alignment = C
        doc.add_heading("1.1 Overview", level=2)
    else:
        _para(doc, "CHAPTER 1", 16, True, C)
        _para(doc, "INTRODUCTION", 16, True, C)
        _para(doc, "1.1 Overview", 14, True, None, spacing=1.5)
    _para(doc, SAMPLE_BODY, 12, False, J, first_indent=0.5)
    if not heading1_style:
        _para(doc, "1.1.1 Scope", 12, True, None)
        _para(doc, "A short sample paragraph.", 12, False, J, first_indent=0.5)
    _para(doc, "Fig 1.1: Sample figure caption", 11, True, C, spacing=1.0)
    _para(doc, "Table 1.1: Sample table caption", 11, True, C, spacing=1.0)
    grid = doc.add_table(rows=2, cols=2)
    grid.style = "Table Grid"
    for r in range(2):
        for c in range(2):
            grid.cell(r, c).text = "Sample"
    _break(doc)
    if not heading1_style:
        _para(doc, "CHAPTER 2", 16, True, C)
        _para(doc, "LITERATURE SURVEY", 16, True, C)
        _para(doc, "Another sample paragraph in the second chapter.", 12, False, J, first_indent=0.5)
        _break(doc)
    _para(doc, "REFERENCES", 16, True, C)
    _para(doc, "[1] A. Author, “Sample reference title,” Sample Journal, vol. 1, no. 1, pp. 1-10, 2020.", 12, False, J, spacing=1.15)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def ppt_template(with_layout_placeholders: bool = True) -> bytes:
    """A deck with a coloured header band, a logo box and a footer on its
    master, a title slide and two sample content slides."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches as In, Pt as P

    prs = Presentation()
    prs.slide_width, prs.slide_height = In(13.333), In(7.5)
    master = prs.slide_master
    # drawn on a scratch slide, then moved onto the master so it shows on every slide
    canvas = prs.slides.add_slide(prs.slide_layouts[6])
    band = canvas.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, In(0.9))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(0x1F, 0x3A, 0x6E)
    band.line.fill.background()
    band.name = "Header band"
    college = canvas.shapes.add_textbox(In(1.3), In(0.15), In(10), In(0.6))
    college.name = "College name"
    r = college.text_frame.paragraphs[0].add_run()
    r.text = "GSSS Institute of Engineering and Technology for Women, Mysuru"
    r.font.size, r.font.bold, r.font.color.rgb = P(20), True, RGBColor(0xFF, 0xFF, 0xFF)
    logo = canvas.shapes.add_shape(MSO_SHAPE.OVAL, In(0.3), In(0.1), In(0.7), In(0.7))
    logo.name = "Logo"
    logo.fill.solid()
    logo.fill.fore_color.rgb = RGBColor(0xF2, 0xB7, 0x05)
    footer = canvas.shapes.add_textbox(In(0.4), In(7.0), In(9), In(0.4))
    footer.name = "Footer"
    fr = footer.text_frame.paragraphs[0].add_run()
    fr.text = "Department of Information Science and Engineering"
    fr.font.size, fr.font.color.rgb = P(12), RGBColor(0x55, 0x55, 0x55)
    for shape in (band, college, logo, footer):
        master.shapes._spTree.append(shape._element)
    ids = prs.slides._sldIdLst
    prs.part.drop_rel(ids[0].rId)
    ids.remove(ids[0])
    # keep the layouts' own title below the band
    for layout in prs.slide_layouts:
        for shape in layout.placeholders:
            if shape.placeholder_format.type is not None and shape.placeholder_format.idx == 0 and layout.name != "Title Slide":
                shape.left, shape.top, shape.width, shape.height = In(0.6), In(1.0), In(12.1), In(1.0)
            if shape.placeholder_format.idx == 1 and layout.name == "Title and Content":
                shape.left, shape.top, shape.width, shape.height = In(0.6), In(2.1), In(12.1), In(4.6)

    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "TITLE OF THE PROJECT"
    title.placeholders[1].text = "Presented by: NAME OF THE STUDENT (USN)\nUnder the guidance of: <<Name of Guide>>\n20XX-20XX"
    if with_layout_placeholders:
        for heading, lines in (("Introduction", ["Sample bullet one", "Sample bullet two"]), ("Objectives", ["Sample objective"])):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = heading
            tf = slide.placeholders[1].text_frame
            tf.text = lines[0]
            for line in lines[1:]:
                tf.add_paragraph().text = line
    else:
        for heading in ("Introduction", "Objectives"):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            box = slide.shapes.add_textbox(In(0.6), In(1.1), In(12), In(0.9))
            hr2 = box.text_frame.paragraphs[0].add_run()
            hr2.text = heading
            hr2.font.size, hr2.font.bold, hr2.font.color.rgb = P(32), True, RGBColor(0x1F, 0x3A, 0x6E)
            body = slide.shapes.add_textbox(In(0.8), In(2.2), In(11.5), In(4.4))
            body.text_frame.word_wrap = True
            br = body.text_frame.paragraphs[0].add_run()
            br.text = "Sample bullet text of the format"
            br.font.size = P(22)
    end = prs.slides.add_slide(prs.slide_layouts[0])
    end.shapes.title.text = "Thank You"
    end.placeholders[1].text = ""
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def png(width=900, height=500, colour=(40, 90, 160)) -> bytes:
    from PIL import Image
    out = io.BytesIO()
    Image.new("RGB", (width, height), colour).save(out, format="PNG")
    return out.getvalue()


DETAILS = {
    "title": "Smart Attendance System Using Face Recognition",
    "students": [{"name": "Asha Rao", "usn": "4GW22IS001"}, {"name": "Bina Kumar", "usn": "4GW22IS002"}, {"name": "Chitra N", "usn": "4GW22IS003"}],
    "guide_name": "Dr. Meera Nair", "guide_designation": "Associate Professor",
    "hod_name": "Dr. Latha Rao", "principal_name": "Dr. Shiva Kumar",
    "department": "Information Science and Engineering", "college": "GSSS Institute of Engineering and Technology for Women",
    "academic_year": "2025-2026", "project_type": "Final year project",
}

CONTENT = {
    "abstract": ["Manual attendance takes class time and is easy to falsify. This project records attendance from a classroom camera using face recognition."],
    "acknowledgement": ["We thank our guide and the Department of Information Science and Engineering for their support."],
    "chapters": [
        {"title": "Introduction", "blocks": [
            {"type": "heading", "level": 2, "text": "Overview"},
            {"type": "para", "text": "Attendance is recorded in every class. The **manual roll call** takes several minutes of each session."},
            {"type": "heading", "level": 2, "text": "Problem Statement"},
            {"type": "para", "text": "To record attendance automatically and reliably without interrupting the class."},
            {"type": "heading", "level": 3, "text": "Scope"},
            {"type": "bullets", "items": ["Single classroom camera", "Students enrolled in advance"]},
        ]},
        {"title": "System Design", "blocks": [
            {"type": "heading", "level": 2, "text": "Architecture"},
            {"type": "para", "text": "The system has a capture module, a recognition module and a records module."},
            {"type": "figure", "image_id": 1, "caption": "System architecture"},
            {"type": "table", "caption": "Modules and their roles", "rows": [["Module", "Role"], ["Capture", "Reads frames from the camera"], ["Recognition", "Matches faces to enrolled students"]]},
            {"type": "figure", "image_id": 2, "caption": "Data flow"},
        ]},
        {"title": "Conclusion", "blocks": [{"type": "para", "text": "The system records attendance without a roll call."}]},
    ],
    "references": ["P. Viola and M. Jones, “Rapid object detection using a boosted cascade of simple features,” CVPR, 2001.",
                   "[2] F. Schroff, D. Kalenichenko and J. Philbin, “FaceNet: A unified embedding for face recognition and clustering,” CVPR, 2015."],
}
