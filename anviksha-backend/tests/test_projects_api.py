"""The project builder through its API: formats, a project from details to
downloads, honesty about gaps, privacy, and guide review. The writing model
is replaced by a stand-in; storage is kept in memory."""
import io
import json

import pytest
from docx import Document
from pptx import Presentation

from app.core import file_storage, project_jobs, project_writer as pw
from app.models import models as m
from tests import project_samples as S
from tests.conftest import auth_headers
from tests.test_missing_endpoints import campus  # noqa: F401  (shared fixture)

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def fake_model(prompt: str, max_tokens: int = 0) -> str:
    """Answers like the real model would, including one made-up figure and a
    citation, so the checks that remove those can be seen working."""
    if "Write the abstract" in prompt:
        return json.dumps({"abstract": "The project records attendance from a classroom camera using face recognition. " * 4 + "It reached 99.9% accuracy."})
    if "presentation slides" in prompt or "Rewrite one presentation slide" in prompt:
        if "Rewrite one" in prompt:
            return json.dumps({"title": "Shorter slide", "bullets": ["One short point", "Second short point"], "notes": "Notes."})
        return json.dumps({"slides": [{"title": "Key points", "bullets": ["First point from the chapter", "Second point from the chapter", "Accuracy was 99.9%"], "notes": "Speaker notes."}]})
    title = prompt.split('write the chapter "')[1].split('"')[0]
    blocks = [{"type": "heading", "level": 2, "text": "1.1 Overview"},
              {"type": "para", "text": f"This section of {title} is written from the student's material. The system was tested on 40 students. It achieved 97.5% accuracy as shown by Smith et al."},
              {"type": "bullets", "items": ["A point from the material", "Reported in [3]"]}]
    if "IMG" in prompt.split("PICTURES THE STUDENT SUPPLIED")[1].split("TABLES FROM")[0]:
        image = prompt.split("PICTURES THE STUDENT SUPPLIED")[1].split("IMG")[1].split(":")[0]
        blocks.append({"type": "figure", "image": f"IMG{image}", "caption": "A picture"})
    if "TBL1" in prompt:
        blocks.append({"type": "table", "table": "TBL1", "caption": "Results table"})
    if "make it shorter" in prompt:
        blocks = [{"type": "para", "text": f"A shorter version of {title} written from the student's material."}]
    return json.dumps({"blocks": blocks, "missing": ["How many classrooms was it tried in?"]})


@pytest.fixture()
def lab(campus, monkeypatch):  # noqa: F811
    store: dict[str, bytes] = {}
    monkeypatch.setattr(file_storage, "upload_file", lambda path, content, content_type=None: store.__setitem__(path, content))
    monkeypatch.setattr(file_storage, "get_file", lambda path: store[path] if path in store else (_ for _ in ()).throw(FileNotFoundError(path)))
    monkeypatch.setattr(file_storage, "delete_file", lambda path: store.pop(path, None))
    monkeypatch.setattr(project_jobs, "RUN_INLINE", True)
    monkeypatch.setattr(project_jobs, "complete_override", fake_model)
    monkeypatch.delenv("PDF_CONVERTER_URL", raising=False)
    monkeypatch.delenv("SOFFICE_PATH", raising=False)
    monkeypatch.setattr("app.core.project_pdf._soffice", lambda: None)
    db = campus["db"]
    admin = m.User(email="admin@test.local", hashed_password="x", role="admin")
    db.add(admin)
    db.commit()
    return {**campus, "store": store, "ah": auth_headers(admin.id)}


def _upload_format(client, headers, kind="report", name="Final year report format", data=None, filename=None, **form):
    data = data if data is not None else (S.report_template() if kind == "report" else S.ppt_template())
    filename = filename or ("format.docx" if kind == "report" else "format.pptx")
    return client.post("/project-formats", headers=headers, data={"kind": kind, "name": name, **form}, files={"file": (filename, io.BytesIO(data), DOCX if kind == "report" else PPTX)})


def _project(client, lab, **extra):
    h = lab["h"]
    report = _upload_format(client, h["hod"]).json()["id"]
    ppt = _upload_format(client, h["hod"], kind="ppt", name="Final year slides format").json()["id"]
    body = {"title": "Smart Attendance System Using Face Recognition", "report_format_id": report, "ppt_format_id": ppt,
            "details": {"guide_name": "Dr. Meera Nair", "guide_designation": "Associate Professor", "hod_name": "Dr. Latha Rao", "principal_name": "Dr. Shiva Kumar",
                        "academic_year": "2025-2026", "project_type": "Final year project", "domain": "Computer vision", "tech_stack": "Python, OpenCV"}, **extra}
    r = client.post("/student/projects", headers=h["s1"], json=body)
    assert r.status_code == 200, r.text
    return r.json()


# ---------------------------------------------------------------- formats

def test_formats_are_checked_on_upload_and_scoped(client, lab):
    h = lab["h"]
    made = _upload_format(client, h["hod"], project_type="Final year project", batch="2023")
    assert made.status_code == 200, made.text
    f = made.json()
    assert (f["scope"], f["kind"], f["analysis"]["page"]["width_in"], f["preview_available"]) == ("official", "report", 8.27, False)
    assert {x["field"] for x in f["field_map"]} >= {"title", "student_name_usn", "guide_name", "academic_year"}
    assert any("format.docx" in k for k in lab["store"])

    assert _upload_format(client, h["hod"], data=b"this is not a word file").status_code == 400            # corrupted
    assert "must be a .docx" in _upload_format(client, h["hod"], filename="format.doc").json()["detail"]
    assert "Save As" in _upload_format(client, h["hod"], filename="old.doc").json()["detail"]
    assert _upload_format(client, h["hod"], kind="ppt", data=S.report_template(), filename="x.pptx").status_code == 400    # a Word file named .pptx
    assert _upload_format(client, h["hod"], name="x").status_code == 400
    assert _upload_format(client, h["t1"]).status_code == 403

    own = _upload_format(client, h["s1"], name="My guide's format").json()
    assert own["scope"] == "custom"
    names = lambda who: sorted(x["name"] for x in client.get("/project-formats", headers=who).json()["formats"])  # noqa: E731
    assert names(h["s1"]) == ["Final year report format", "My guide's format"]
    assert names(h["s2"]) == ["Final year report format"]                                                    # another student's own format is private
    assert names(h["hod_cse"]) == []                                                # another department
    assert client.get(f"/project-formats/{own['id']}", headers=h["s2"]).status_code == 404

    # the uploader corrects the field map; nobody else may
    fixed = client.patch(f"/project-formats/{f['id']}", headers=h["hod"], json={"field_map": [{"match": "TITLE OF THE PROJECT", "field": "title"}, {"match": "zz", "field": "bogus"}]})
    assert fixed.json()["field_map"] == [{"match": "TITLE OF THE PROJECT", "field": "title"}]
    assert client.patch(f"/project-formats/{f['id']}", headers=h["s1"], json={"name": "Hacked name"}).status_code == 403
    assert client.patch(f"/project-formats/{f['id']}", headers=h["hod"], json={"active": False}).json()["active"] is False
    assert names(h["s2"]) == []                                                                              # switched off: no longer offered
    assert client.get(f"/project-formats/{f['id']}/file?as_pdf=true", headers=h["hod"]).status_code == 503   # no converter in tests
    assert client.get(f"/project-formats/{f['id']}/file", headers=h["hod"]).content[:2] == b"PK"

    # deleting: only the uploader, and never a format a project depends on
    assert client.delete(f"/project-formats/{f['id']}", headers=h["s1"]).status_code in (403, 404)
    used = client.post("/student/projects", headers=h["s1"], json={"title": "A project using my format", "report_format_id": own["id"]}).json()
    assert client.delete(f"/project-formats/{own['id']}", headers=h["s1"]).status_code == 409
    client.delete(f"/student/projects/{used['id']}", headers=h["s1"])
    assert client.delete(f"/project-formats/{own['id']}", headers=h["s1"]).status_code == 200
    assert client.delete(f"/project-formats/{f['id']}", headers=h["hod"]).status_code == 200
    assert not any(k.startswith("project_formats/") and ("My_guide" in k or f"{f['id']}_" in k) for k in lab["store"])


# ---------------------------------------------------------------- a project, end to end

def test_details_are_prefilled_and_teammates_come_from_the_records(client, lab):
    h = lab["h"]
    info = client.get("/student/projects/start", headers=h["s1"]).json()
    assert info["me"] == {"name": "Asha Rao", "usn": "T23IS001"} and info["department"] == "ISE (TEST)" and info["academic_year"] == "2026-2027"
    assert {g["name"] for g in info["guides"]} == {"Teacher One", "Teacher Two"}
    assert client.get("/student/projects/classmates?q=bina", headers=h["s1"]).json() == [{"name": "Bina Kumar", "usn": "T23IS002"}]
    assert client.get("/student/projects/classmates?q=chitra", headers=h["s1"]).json() == []                 # another batch
    assert client.get("/student/projects/classmates?q=b", headers=h["s1"]).json() == []
    p = _project(client, lab)
    assert p["details"]["students"] == [{"name": "Asha Rao", "usn": "T23IS001"}]                              # the creator is on the team by default
    assert client.post("/student/projects", headers=h["s1"], json={"title": "abc"}).status_code == 400


def test_it_asks_for_what_is_missing_instead_of_inventing(client, lab):
    h = lab["h"]
    p = _project(client, lab)
    pid = p["id"]
    assert {c["status"] for c in p["readiness"]} == {"missing"}
    r = client.post(f"/student/projects/{pid}/jobs", headers=h["s1"], json={"kind": "generate"})
    assert r.status_code == 400 and "no material" in r.json()["detail"]

    client.patch(f"/student/projects/{pid}", headers=h["s1"], json={"answers": {"problem_statement": "Manual roll call wastes class time and can be falsified.",
                                                                                 "objectives": "Record attendance automatically.\nKeep a record per subject."}})
    state = client.get(f"/student/projects/{pid}", headers=h["s1"]).json()
    intro = state["readiness"][0]
    assert (intro["status"], intro["missing"]) == ("ready", []) and [q["answered"] for q in intro["questions"]] == [True, True, False, False]
    r = client.post(f"/student/projects/{pid}/jobs", headers=h["s1"], json={"kind": "generate"})
    assert r.status_code == 422 and "Testing" in r.json()["detail"]["chapters"] and "Introduction" not in r.json()["detail"]["chapters"]

    job = client.post(f"/student/projects/{pid}/jobs", headers=h["s1"], json={"kind": "generate", "proceed_with_gaps": True}).json()
    assert job["status"] == "done", job
    content = client.get(f"/student/projects/{pid}", headers=h["s1"]).json()["content"]
    written = {c["title"]: c for c in content["chapters"]}
    assert written["Introduction"]["blocks"] and written["Testing"]["blocks"] == [] and written["Testing"]["skipped"] is True
    assert "Tests you ran" in written["Testing"]["missing"]
    assert content["references"] == []                                                                       # none given, none made up


def test_full_flow_builds_files_in_the_format_and_keeps_versions(client, lab):
    h, store = lab["h"], lab["store"]
    p = _project(client, lab, details={"students": [{"name": "Asha Rao", "usn": "T23IS001"}, {"name": "Bina Kumar", "usn": "T23IS002"}], "guide_name": "Dr. Meera Nair",
                                       "guide_designation": "Associate Professor", "hod_name": "Dr. Latha Rao", "principal_name": "Dr. Shiva Kumar", "academic_year": "2025-2026"})
    pid = p["id"]
    notes = ("Introduction and overview. The problem is that manual roll call wastes class time. The objective is automatic attendance.\n\n" * 6
             + "System design and architecture. The modules are capture, recognition and records. The system was tested on 40 students.\n\n" * 6
             + "Testing. Test case 1: known face, expected present, actual present.\n\nResults. The output was observed on 40 students.\n\n" * 6
             + "Implementation. The code was developed in Python using the OpenCV library.\n\nConclusion and future scope: add liveness detection.\n\n" * 6
             + "Literature survey of related work and existing papers studied.\n\nSoftware requirement: Python. Hardware requirement: a camera.\n\n" * 6)
    up = client.post(f"/student/projects/{pid}/materials", headers=h["s1"], files={"file": ("notes.txt", notes.encode(), "text/plain")}).json()
    assert up["kind"] == "text" and up["characters"] > 1000
    assert client.post(f"/student/projects/{pid}/materials", headers=h["s1"], files={"file": ("shot.png", S.png(), "image/png")}).json()["pictures"] == 1
    assert client.post(f"/student/projects/{pid}/materials", headers=h["s1"], files={"file": ("results.csv", b"Metric,Value\nStudents,40\nSessions,12\n", "text/csv")}).json()["tables"] == 1
    assert client.post(f"/student/projects/{pid}/materials", headers=h["s1"], files={"file": ("virus.exe", b"MZ", "application/octet-stream")}).status_code == 400
    assert client.post(f"/student/projects/{pid}/materials", headers=h["s1"], files={"file": ("broken.pdf", b"not a pdf", "application/pdf")}).status_code == 400
    client.patch(f"/student/projects/{pid}", headers=h["s1"], json={"answers": {"modules": "Capture - reads frames\nRecognition - matches faces\nRecords - stores attendance",
                                                                                 "references": "[1] P. Viola and M. Jones, Rapid object detection, CVPR 2001.\nshort"}})
    state = client.get(f"/student/projects/{pid}", headers=h["s1"]).json()
    assert all(c["status"] == "ready" for c in state["readiness"]), [(c["title"], c["status"]) for c in state["readiness"]]

    job = client.post(f"/student/projects/{pid}/jobs", headers=h["s1"], json={"kind": "generate"}).json()
    assert job["status"] == "done" and job["result"] == {"version": 1, "files": ["ppt_pptx", "report_docx"]}, job
    state = client.get(f"/student/projects/{pid}", headers=h["s1"]).json()
    content = state["content"]
    text = json.dumps({k: content[k] for k in ("chapters", "slides", "abstract")})
    assert "97.5" not in text and "99.9" not in text and "et al" not in text and "[3]" not in text           # invented figure and citations removed
    assert "tested on 40 students" in text                                                                    # a figure the student did give is kept
    assert sum("is not in your material" in w for w in content["warnings"]) >= 8
    assert content["references"] == ["P. Viola and M. Jones, Rapid object detection, CVPR 2001."]
    assert any(i["source"] == "generated" for i in state["images"])                                           # the module diagram, labelled as drawn
    figures = [b["image_id"] for c in content["chapters"] for b in c["blocks"] if b["type"] == "figure"]
    assert len(figures) == len(set(figures)) >= 1                                                             # each picture is placed once
    with_pictures = {c["key"] for c in content["chapters"] if any(b["type"] == "figure" for b in c["blocks"])}
    assert with_pictures <= {"design", "implementation", "testing", "results"}                                # never dropped into the introduction
    tables = [b for c in content["chapters"] for b in c["blocks"] if b["type"] == "table"]
    assert len(tables) == 1 and "source_table" not in tables[0]                                               # the student's table appears once
    table = tables[0]
    assert table["rows"] == [["Metric", "Value"], ["Students", "40"], ["Sessions", "12"]]                     # rows copied from the student's file
    assert content["slides"][0]["title"] == "Agenda" and content["slides"][-1]["title"] == "References"

    v = state["versions"][0]
    assert v["number"] == 1 and v["report"]["report"]["unfilled"] == [] and v["report"]["report"]["leftover_wording"] == [] and v["report"]["slides"]["too_much_text_on"] == []
    assert any("PDF export is not set up" in n for n in v["report"]["notes"])
    docx = client.get(f"/student/projects/{pid}/versions/1/report_docx", headers=h["s1"])
    assert docx.status_code == 200 and "Smart_Attendance" in docx.headers["content-disposition"]
    words = "\n".join(p.text for p in Document(io.BytesIO(docx.content)).paragraphs)
    assert "ASHA RAO (T23IS001)" in words and "BINA KUMAR (T23IS002)" in words and "CHAPTER 8" in words and "Dr. Meera Nair" in words and "sample" not in words.lower()
    deck = Presentation(io.BytesIO(client.get(f"/student/projects/{pid}/versions/1/ppt_pptx", headers=h["s1"]).content))
    assert deck.slides[0].shapes.title.text == "SMART ATTENDANCE SYSTEM USING FACE RECOGNITION" and len(deck.slides) >= 10
    assert client.get(f"/student/projects/{pid}/versions/1/report_pdf", headers=h["s1"]).status_code == 404  # honest: no PDF was made
    assert client.get(f"/student/projects/{pid}/versions/1/secrets", headers=h["s1"]).status_code == 404

    # edit by hand, rewrite one chapter and one slide with an instruction, rebuild: a new version, the old one kept
    content["chapters"][0]["blocks"].insert(0, {"type": "para", "text": "A sentence the student typed by hand."})
    content["slides"] = [content["slides"][1], content["slides"][0]] + content["slides"][2:]                   # reorder
    content["chapters"][0]["blocks"].append({"type": "figure", "image_id": 99999, "caption": "someone else's picture"})
    saved = client.patch(f"/student/projects/{pid}", headers=h["s1"], json={"content": content}).json()["content"]
    assert saved["chapters"][0]["blocks"][0]["text"] == "A sentence the student typed by hand." and saved["slides"][1]["title"] == "Agenda"
    assert not any(b.get("image_id") == 99999 for b in saved["chapters"][0]["blocks"])                        # a picture that is not this project's is refused
    key = saved["chapters"][2]["key"]
    assert client.post(f"/student/projects/{pid}/jobs", headers=h["s1"], json={"kind": "rewrite", "target": "chapter", "key": key, "instruction": "make it shorter"}).json()["status"] == "done"
    assert client.post(f"/student/projects/{pid}/jobs", headers=h["s1"], json={"kind": "rewrite", "target": "slide", "index": 2}).json()["status"] == "done"
    state = client.get(f"/student/projects/{pid}", headers=h["s1"]).json()
    assert state["content"]["chapters"][2]["blocks"][0]["text"].startswith("A shorter version") and state["content"]["slides"][2]["title"] == "Shorter slide"
    assert state["content"]["chapters"][0]["blocks"][0]["text"] == "A sentence the student typed by hand."     # other chapters untouched
    assert client.post(f"/student/projects/{pid}/jobs", headers=h["s1"], json={"kind": "build"}).json()["result"]["version"] == 2
    state = client.get(f"/student/projects/{pid}", headers=h["s1"]).json()
    assert [v["number"] for v in state["versions"]] == [2, 1]
    assert client.get(f"/student/projects/{pid}/versions/1/report_docx", headers=h["s1"]).content == docx.content
    restored = client.post(f"/student/projects/{pid}/versions/1/restore", headers=h["s1"]).json()
    assert restored["content"]["slides"][0]["title"] == "Agenda" and len(restored["versions"]) == 2

    # pictures: caption, then remove everywhere
    image_id = figures[0]
    assert client.patch(f"/student/projects/{pid}/images/{image_id}", headers=h["s1"], json={"caption": "Login screen"}).json()["caption"] == "Login screen"
    assert client.get(f"/student/projects/{pid}/images/{image_id}", headers=h["s1"]).content[:4] in (b"\x89PNG", b"\xff\xd8\xff\xe0")
    client.delete(f"/student/projects/{pid}/images/{image_id}", headers=h["s1"])
    after = client.get(f"/student/projects/{pid}", headers=h["s1"]).json()
    assert image_id not in [b.get("image_id") for c in after["content"]["chapters"] for b in c["blocks"]] and image_id not in [i["id"] for i in after["images"]]

    # privacy: nobody else reaches the project or its files
    for other in (h["s2"], h["t1"], h["hod"], lab["ah"]):
        assert client.get(f"/student/projects/{pid}", headers=other).status_code in (403, 404)
        assert client.get(f"/student/projects/{pid}/versions/1/report_docx", headers=other).status_code in (403, 404)
    assert client.get("/student/projects", headers=h["s2"]).json() == []

    assert client.delete(f"/student/projects/{pid}", headers=h["s1"]).status_code == 200
    assert not any(k.startswith(f"projects/{pid}/") for k in store)                                            # its files are gone from storage too


def test_guide_review_only_after_sharing(client, lab):
    db, h = lab["db"], lab["h"]
    p = _project(client, lab)
    pid = p["id"]
    client.patch(f"/student/projects/{pid}", headers=h["s1"], json={"answers": {"problem_statement": "Manual roll call wastes class time and is easy to falsify.", "objectives": "Automatic attendance for every class."}})
    client.post(f"/student/projects/{pid}/jobs", headers=h["s1"], json={"kind": "generate", "proceed_with_gaps": True})
    t1, t2 = lab["t1"], lab["t2"]
    assert client.get("/teacher/project-reviews", headers=h["t1"]).json() == []
    assert client.get(f"/teacher/project-reviews/{pid}", headers=h["t1"]).status_code == 404
    assert client.put(f"/student/projects/{pid}/share", headers=h["s1"], json={"faculty_id": lab["t3"].id if "t3" in lab else 99999}).status_code == 404   # not in the department
    shared = client.put(f"/student/projects/{pid}/share", headers=h["s1"], json={"faculty_id": t1.id}).json()
    assert shared["shared_with"] == {"id": t1.id, "name": "Teacher One"}
    login = db.query(m.User).filter(m.User.faculty_id == t1.id).one()
    assert any("shared the project" in n.message for n in db.query(m.Notification).filter(m.Notification.user_id == login.id))

    listing = client.get("/teacher/project-reviews", headers=h["t1"]).json()
    assert [(x["title"], x["latest_version"]) for x in listing] == [("Smart Attendance System Using Face Recognition", 1)]
    seen = client.get(f"/teacher/project-reviews/{pid}", headers=h["t1"]).json()
    assert seen["content"]["chapters"][0]["blocks"] and "answers" not in seen and "materials" not in seen       # the draft, not the raw uploads
    assert client.get(f"/teacher/project-reviews/{pid}/versions/1/report_docx", headers=h["t1"]).content[:2] == b"PK"
    assert client.get(f"/teacher/project-reviews/{pid}", headers=h["t2"]).status_code == 404                    # another teacher
    c = client.post(f"/teacher/project-reviews/{pid}/comments", headers=h["t1"], json={"section": "introduction", "body": "Tighten the problem statement."}).json()
    assert (c["author"], c["role"], c["resolved"]) == ("Teacher One", "teacher", False)
    assert client.post(f"/teacher/project-reviews/{pid}/comments", headers=h["t2"], json={"section": "x", "body": "hello"}).status_code == 404
    mine = client.get(f"/student/projects/{pid}", headers=h["s1"]).json()
    assert [x["body"] for x in mine["comments"]] == ["Tighten the problem statement."]
    assert client.patch(f"/student/projects/{pid}/comments/{c['id']}", headers=h["s1"]).json()["resolved"] is True
    client.post(f"/student/projects/{pid}/comments", headers=h["s1"], json={"section": "introduction", "body": "Done, please check."})
    assert client.put(f"/student/projects/{pid}/share", headers=h["s1"], json={"faculty_id": None}).json()["shared_with"] is None
    assert client.get(f"/teacher/project-reviews/{pid}", headers=h["t1"]).status_code == 404                    # unshared: access ends
    del t2


def test_writer_rules_hold_without_the_model():
    assert pw.strip_unsupported("It works. Accuracy was 97.5%. It was tried on 40 students.", pw._digits("tried on 40 students")) == ("It works. It was tried on 40 students.", ["Accuracy was 97.5%."])
    assert pw.strip_unsupported("Speed rose by 12.50 %.", pw._digits("12.5%"))[1] == []                         # the same figure written differently is recognised
    assert pw.clean_references("1. Short\n[2] A. Author, A real long reference title, 2020.\n[2] A. Author, A real long reference title, 2020.") == ["A. Author, A real long reference title, 2020."]
    assert pw.parse_modules("1. Capture - reads frames\n- Recognition: matches faces\nRecords") == ["Capture", "Recognition", "Records"]
    assert pw.module_diagram(["Only one"]) is None and pw.module_diagram(["A", "B", "C"])[:4] == b"\x89PNG"
    plan = pw.chapter_plan(["Introduction", "Literature Review", "Proposed Methodology", "Implementation", "Results", "Conclusion"])
    assert [c["key"] for c in plan] == ["introduction", "literature", "other:proposed-methodology", "implementation", "results", "conclusion"]
    assert [c["key"] for c in pw.chapter_plan(["Introduction"])][:2] == ["introduction", "literature"]           # too few in the format: the usual eight
    lines = pw.acknowledgement({"guide_name": "Dr. A", "college": "Test College"})
    assert any("Dr. A" in line for line in lines) and not any("Principal" in line or "Head of the Department" in line for line in lines)
    out = pw.write_chapter({"key": "testing", "title": "Testing"}, {}, {}, [], [], [], lambda *a, **k: (_ for _ in ()).throw(AssertionError("the model must not be asked")))
    assert out["skipped"] is True and out["blocks"] == []
