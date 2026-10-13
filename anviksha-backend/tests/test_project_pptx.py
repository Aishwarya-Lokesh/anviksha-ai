"""The presentation engine: reading a format and writing slides into it."""
import io

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.core import project_pptx as pp
from tests import project_samples as S

SLIDES = [
    {"title": "Introduction", "bullets": ["Attendance is taken in every class", "A roll call takes several minutes"], "notes": "Open with the problem."},
    {"title": "Architecture", "bullets": ["Capture module", "Recognition module", "Records module"], "image_id": 1},
    {"title": "Results", "bullets": [f"Observation number {i} recorded during testing of the system" for i in range(1, 15)]},
]


def _deck(data: bytes):
    prs = Presentation(io.BytesIO(data))
    out = []
    for slide in prs.slides:
        r = pp.roles(slide)
        out.append((pp._text(r["title"]) if r["title"] is not None else "", [p.text for p in r["body"].text_frame.paragraphs] if r["body"] is not None else []))
    return prs, out


def test_analysis_reads_the_format():
    a = pp.analyse(S.ppt_template())
    assert a["slide"] == {"width_in": 13.33, "height_in": 7.5, "shape": "widescreen 16:9"}
    assert [(b["name"], b["kind"]) for b in a["on_every_slide"]] == [("Header band", "shape"), ("College name", "text"), ("Logo", "shape"), ("Footer", "text")]
    assert [s["heading"] for s in a["sample_slides"]] == ["TITLE OF THE PROJECT", "Introduction", "Objectives", "Thank You"]
    assert a["content_slides_use"] == "the format's own layout" and a["warnings"] == []
    assert {m["match"]: m["field"] for m in a["field_map"]} == {"TITLE OF THE PROJECT": "title", "NAME OF THE STUDENT (USN)": "student_name_usn", "<<Name of Guide>>": "guide_name", "20XX-20XX": "academic_year"}
    assert pp.analyse(S.ppt_template(with_layout_placeholders=False))["content_slides_use"] == "a copy of the format's sample slide"
    with pytest.raises(ValueError):
        pp.analyse(b"nope")


@pytest.mark.parametrize("placeholders", [True, False])
def test_deck_keeps_the_format_and_carries_the_students_slides(placeholders):
    template = S.ppt_template(with_layout_placeholders=placeholders)
    a = pp.analyse(template)
    out = pp.build(template, a["field_map"], S.DETAILS, SLIDES, {1: S.png()})
    prs, deck = _deck(out["pptx"])
    before = Presentation(io.BytesIO(template))

    # the format itself: size, master artwork and layouts are as they were
    assert (prs.slide_width, prs.slide_height) == (before.slide_width, before.slide_height)
    assert [(s.name, s.left, s.top, s.width, s.height) for s in prs.slide_master.shapes] == [(s.name, s.left, s.top, s.width, s.height) for s in before.slide_master.shapes]
    assert [l.name for l in prs.slide_layouts] == [l.name for l in before.slide_layouts]

    titles = [t for t, _ in deck]
    assert titles[0] == "SMART ATTENDANCE SYSTEM USING FACE RECOGNITION" and titles[-1] == "Thank You"
    assert titles[1:3] == ["Introduction", "Architecture"] and titles[3] == "Results" and titles[4].startswith("Results (contd.)")
    assert "Objectives" not in titles and not any("Sample" in b for _, bullets in deck for b in bullets)        # the format's sample slides are gone
    assert deck[1][1] == ["Attendance is taken in every class", "A roll call takes several minutes"]
    first = prs.slides[0]
    assert "Presented by: ASHA RAO (4GW22IS001), BINA KUMAR (4GW22IS002) and CHITRA N (4GW22IS003)" in " ".join(sh.text_frame.text for sh in first.shapes if sh.has_text_frame)
    assert "Dr. Meera Nair" in " ".join(sh.text_frame.text for sh in first.shapes if sh.has_text_frame)
    assert out["unfilled"] == [] and pp.leftovers(out["pptx"], a["field_map"], ["Sample bullet", "Sample objective"]) == []

    # every slide's text fits its box at the format's own font size
    assert out["overflow"] == [] and out["continued"] >= 1
    assert all(s["lines_used"] <= s["lines_available"] for s in out["slides"])
    results = [bullets for t, bullets in deck if t.startswith("Results")]
    assert sum(len(b) for b in results) == 14 and [b for part in results for b in part] == SLIDES[2]["bullets"]   # nothing lost or reordered

    # the picture sits beside the text, inside the slide, not over it
    arch = prs.slides[2]
    picture = next(sh for sh in arch.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    body = pp.roles(arch)["body"]
    assert picture.left >= body.left + body.width and picture.left + picture.width <= prs.slide_width and picture.top + picture.height <= prs.slide_height
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Open with the problem."

    # body text is written at the size the format uses, never shrunk
    sample_size = pp._inherited_pt(before.slides[1], pp.roles(before.slides[1])["body"], "body")
    assert {s["font_pt"] for s in out["slides"]} == {sample_size}


def test_long_bullets_are_split_at_sentences_and_missing_pictures_reported():
    assert pp.split_bullets(["One. Two. Three."], lines=6, chars=40) == [["One. Two. Three."]]
    long = " ".join(f"Sentence number {i} explains one more detail of the design." for i in range(1, 13))
    pages = pp.split_bullets([long], lines=5, chars=40)
    assert len(pages) > 1 and " ".join(b for page in pages for b in page) == long
    assert all(sum(pp._lines_needed(b, 40) for b in page) <= 5 for page in pages)
    assert pp.split_bullets([], 5, 40) == [[]]

    a = pp.analyse(S.ppt_template())
    out = pp.build(S.ppt_template(), a["field_map"], {"title": "T", "students": []}, [{"title": "Design", "bullets": ["x"], "image_id": 9}], {})
    assert out["dropped"] == ["Design"] and "Student name with USN (one per student)" in out["unfilled"]
    prs, deck = _deck(out["pptx"])
    assert not any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in prs.slides[1].shapes)
